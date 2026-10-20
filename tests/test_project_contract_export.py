"""独立项目口径、同源文件与不足数量必须保持一致。"""
import asyncio
import json
from datetime import timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.conductor.nodes import output, planner
from src.conductor.task_spec import TaskSpec


QUERY = ("从国家电网、南方电网及其省级公司官网，查找近90天配电网自动化、智能巡检相关项目，"
         "目标8个独立项目。同一项目的多份公告归并，保留各公告链接。整理项目名称、业主、地区、"
         "建设内容、金额及单位、发布时间、当前阶段和原文依据，输出Excel、JSON和一页分析报告。"
         "不足8个如实说明，未披露字段不得猜测。")


def test_old_spec_dump_omits_optional_project_profile():
    spec = TaskSpec(intent="旧任务")
    assert "project_collection" not in spec.model_dump(mode="json")


def test_project_contract_accepts_excel_and_rejects_conflicting_profiles():
    spec = TaskSpec.from_draft({"intent": "项目", "outputs": ["xlsx", "json"],
                              "project_collection": {"target_count": 8, "topics": ["智能巡检"]}})
    assert spec.project_collection.target_count == 8
    assert spec.project_collection.project_kinds == []
    assert spec.outputs == ["xlsx", "json"]
    with pytest.raises(ValueError):
        TaskSpec.from_draft({"project_collection": {"target_count": 8}, "evidence_collection": {
            "source_evidence": "小红书", "queries": [{"name": "项目", "keyword": "项目"}], "raw_only": True}})


def test_project_planner_preserves_user_quantity_dates_outputs_and_scope(monkeypatch):
    async def plan(*args, **kwargs):
        return {"intent": "仅招标项目", "site_domains": ["sgcc.com.cn", "csg.cn"],
                "keywords": ["智能巡检 招标公告"], "max_items": 6, "data_type": "bid",
                "analysis_instruction": "只接受招标公告", "outputs": ["report_md", "json"]}, None

    monkeypatch.setattr(planner, "_plan", plan)
    monkeypatch.setattr(planner, "lesson_for_planner", lambda *args, **kwargs: "")
    result = asyncio.run(planner.planner_node({"user_input": QUERY}))
    spec = result["task_spec"]
    scope = spec.project_collection
    assert result["needs_clarification"] is False
    assert spec.max_items == scope.target_count == 8
    assert set(spec.outputs) == {"xlsx", "json", "report_md"}
    assert spec.data_type == "generic" and scope.project_kinds == []
    assert spec.analysis_instruction == QUERY
    assert scope.topics == ["配电网自动化", "智能巡检"]
    assert [query.keyword for query in scope.queries] == ["配电网自动化 项目", "智能巡检 项目"]
    assert scope.allow_shortfall is True
    assert scope.publication_to - scope.publication_from == timedelta(days=90)
    assert scope.frozen_at.tzinfo is not None


def _state():
    spec = TaskSpec.from_draft({"intent": "项目核对", "max_items": 8,
        "outputs": ["xlsx", "json", "report_md"], "project_collection": {"target_count": 8}})
    source = {"source_id": 1, "url": "https://official.example/project", "title": "项目公告",
              "content": "智能巡检建设项目，业主甲方。" + "中文" * 20000}
    return {"task_id": "project-export", "task_spec": spec, "collector_used": "search",
        "cleaned_dataset": [source], "analysis": "仅找到1个项目，不能证明达到8个。",
        "quality": {"passed": False, "review_complete": True, "score": 35,
                    "source_checks": [{"source_id": 1, "usable": True, "reason": "原文符合主题"}]},
        "project_result": {"projects": [{"project_id": "p1", "project_name": "=1+1", "owner": "甲方",
            "region": None, "construction_content": "智能巡检建设", "amount": None, "amount_unit": None,
            "publication_date": "2026-10-01", "current_stage": "建设公告", "source_ids": [1],
            "evidence": [{"source_id": 1, "field": "construction_content", "quote": "智能巡检建设项目"}],
            "missing_fields": ["region", "amount", "amount_unit"]}], "sources": [source],
            "excluded": [], "coverage": {"complete": True, "verified_projects": 8, "stop_reason": "budget"}},
        "trace": [{"node": "collect", "ms": 120}]}


def test_project_exports_real_excel_and_full_sources_with_honest_shortfall(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    state["project_result"]["projects"][0]["field_review_notes"] = {"region": "来源1：提案值缺少逐字依据，未采纳并留空。"}
    stage_note = "来源1于2026-10-01披露阶段，仅反映原文时点，冻结日当前进展未另行确认。"
    state["project_result"]["projects"][0]["stage_note"] = stage_note
    result = asyncio.run(output.output_node(state))
    files = result["outputs"]
    assert {"json", "sources_file", "xlsx", "report_md", "trace_file"} <= set(files)
    from src.api.routes.chat import _FILE_MAP
    assert _FILE_MAP["sources_file"][0] == "sources.json"
    from src.api.routes.chat import _build_result
    api = _build_result("test-owner", "test-conversation", {**state, **result}, result["reply"], None, None, QUERY)
    assert any(file["name"] == "sources.json" for file in api["files"])
    data = json.loads(Path(files["json"]).read_text(encoding="utf-8"))
    assert len(data["projects"]) == 1
    assert data["coverage"]["verified_projects"] == 1
    assert data["coverage"]["shortfall"] == 7 and not data["coverage"]["complete"]
    assert data["projects"][0]["amount"] is None
    assert data["projects"][0]["field_review_notes"] == state["project_result"]["projects"][0]["field_review_notes"]
    assert data["projects"][0]["stage_note"] == stage_note
    assert data["projects"][0]["source_links"] == ["https://official.example/project"]
    sources = json.loads(Path(files["sources_file"]).read_text(encoding="utf-8"))
    assert sources[0]["content"] == state["project_result"]["sources"][0]["content"]
    report = Path(files["report_md"]).read_text(encoding="utf-8")
    assert "1/8" in report and "缺少 7" in report and "未披露" in report
    assert result["error"] and result["grade"]["items"] == 1
    workbook = load_workbook(files["xlsx"], read_only=True)
    try:
        assert {"项目清单", "来源核对", "证据明细", "缺口说明"} <= set(workbook.sheetnames)
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
        assert "=1+1" in list(workbook["项目清单"].values)[1]
        rows = list(workbook["项目清单"].values)
        assert json.loads(rows[1][rows[0].index("字段核对说明")]) == data["projects"][0]["field_review_notes"]
        assert rows[1][rows[0].index("阶段核对说明")] == stage_note
        assert rows[1][rows[0].index("当前阶段")] == data["projects"][0]["current_stage"]
    finally:
        workbook.close()


@pytest.mark.parametrize("historical", [False, True])
def test_report_counts_null_fields_without_requiring_notes_and_keeps_recheck_scope(tmp_path, monkeypatch, historical):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    state["project_result"]["projects"][0]["current_stage"] = "试运行" if historical else None
    if historical:
        state["project_result"]["projects"][0]["stage_note"] = "来源2026-08-19披露试运行，冻结日进展未另行确认。"
    state["project_result"]["coverage"]["review_mode_note"] = "仅复核保存的1份原件，未重新搜索。"
    result = asyncio.run(output.output_node(state))
    report = Path(result["outputs"]["report_md"]).read_text(encoding="utf-8")
    assert "项目级金额未确认 1 个" in report
    assert "当前阶段待核实 1 个" in report
    assert "仅复核保存的1份原件，未重新搜索。" in report


def test_empty_project_result_is_diagnostic_not_success(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    state["project_result"]["projects"] = []
    state["quality"] = {"passed": True, "review_complete": True, "score": 90}
    result = asyncio.run(output.output_node(state))
    assert result["error"] and "0/8" in result["reply"]
    assert result["grade"]["items"] == 0 and result["grade"]["level"] == "未完成"
    assert Path(result["outputs"]["xlsx"]).is_file()
    data = json.loads(Path(result["outputs"]["json"]).read_text(encoding="utf-8"))
    assert data["quality"]["passed"] is False and not data["coverage"]["result_verified"]


def test_allowed_shortfall_can_verify_results_without_claiming_target(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    state["task_spec"].project_collection.allow_shortfall = True
    state["quality"].update(passed=True, review_complete=True, score=90)
    result = asyncio.run(output.output_node(state))
    assert result["error"] is None and result["quality"]["passed"]
    coverage = result["project_result"]["coverage"]
    assert not coverage["target_met"] and not coverage["complete"] and coverage["shortfall"] == 7
    assert result["grade"]["level"] == "部分结果已核验" and "不能称目标数量已达成" in result["reply"]


def test_generic_excel_request_exports_without_project_profile(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = {"task_id": "generic-xlsx", "task_spec": TaskSpec(intent="资料", outputs=["xlsx", "json"]),
             "cleaned_dataset": [{"title": "=1+1", "content": "原始资料", "url": "https://example.com"}]}
    result = asyncio.run(output.output_node(state))
    workbook = load_workbook(result["outputs"]["xlsx"], read_only=True)
    try:
        assert list(workbook["采集资料"].values)[1][0] == "=1+1"
        assert all(cell.data_type != "f" for row in workbook["采集资料"] for cell in row)
    finally:
        workbook.close()


def test_project_profile_cannot_be_injected_into_unrelated_task(monkeypatch):
    assert planner._source_conflict({"project_collection": {"target_count": 8}}, "分析新闻口碑")


def test_conflicting_disclosed_values_are_not_labeled_undisclosed(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    state["project_result"]["projects"][0]["field_variants"] = {"amount": ["100", "200"]}
    result = asyncio.run(output.output_node(state))
    data = json.loads(Path(result["outputs"]["json"]).read_text(encoding="utf-8"))
    assert data["projects"][0]["amount"] is None
    assert "多公告字段差异" in Path(result["outputs"]["report_md"]).read_text(encoding="utf-8")
    workbook = load_workbook(result["outputs"]["xlsx"], read_only=True)
    try:
        rows = list(workbook["项目清单"].values)
        assert rows[1][rows[0].index("金额")] == "多公告存在差异，见差异列"
    finally:
        workbook.close()


def test_planner_uses_supported_fields_instead_of_model_invented_aliases(monkeypatch):
    async def plan(*args, **kwargs):
        return {"project_collection": {"target_count": 8, "required_fields": ["project_name", "owner",
            "region", "construction_content", "amount", "currency_unit", "publish_date", "current_stage",
            "source_evidence", "announcement_links"]}, "keywords": ["智能巡检"]}, None

    monkeypatch.setattr(planner, "_plan", plan)
    monkeypatch.setattr(planner, "lesson_for_planner", lambda *args, **kwargs: "")
    result = asyncio.run(planner.planner_node({"user_input": QUERY}))
    assert result["task_spec"].project_collection.required_fields == ["project_name", "owner", "region",
        "construction_content", "amount", "amount_unit", "publication_date", "current_stage"]


def test_one_page_report_has_index_and_summary_without_long_field_table(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    detail = "智能巡检建设详细原文" * 200
    state["project_result"]["projects"][0]["construction_content"] = detail
    result = asyncio.run(output.output_node(state))
    report = Path(result["outputs"]["report_md"]).read_text(encoding="utf-8")
    assert detail not in report
    assert "项目索引" in report and "https://official.example/project" in report
    assert "建设内容 | 金额及单位" not in report
    data = json.loads(Path(result["outputs"]["json"]).read_text(encoding="utf-8"))
    assert data["projects"][0]["construction_content"] == detail


@pytest.mark.parametrize("status", ["unknown", "list_excerpt"])
def test_pending_search_clue_is_not_presented_as_complete_original(tmp_path, monkeypatch, status):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    state["project_result"]["sources"][0]["metadata"] = {"body_status": status}
    result = asyncio.run(output.output_node(state))
    report = Path(result["outputs"]["report_md"]).read_text(encoding="utf-8")
    assert "已取得的正文、搜索线索及状态随 sources.json 保留" in report
    assert "完整来源原件" not in report
    workbook = load_workbook(result["outputs"]["xlsx"], read_only=True)
    try:
        rows = list(workbook["来源核对"].values)
        assert "已取得的来源内容" in rows[0] and "正文取得状态" in rows[0]
        assert rows[1][rows[0].index("正文取得状态")] == status
        assert "仅为线索" in rows[1][rows[0].index("review_scope")]
    finally:
        workbook.close()


def test_unknown_total_and_expired_stage_keep_precise_notes_in_files(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "PROJECT_ROOT", tmp_path)
    state = _state()
    project = state["project_result"]["projects"][0]
    project.update(amount=None, amount_unit="万元", current_stage=None, announcement_count=1,
        amount_scope_note="项目级金额未确认，不代表标包金额未披露",
        stage_note="投标截止日已过，当前阶段待核实")
    result = asyncio.run(output.output_node(state))
    data = json.loads(Path(result["outputs"]["json"]).read_text(encoding="utf-8"))
    exported = data["projects"][0]
    assert exported["amount"] is None and exported["amount_unit"] == "万元"
    assert exported["current_stage"] is None and exported["stage_note"] == project["stage_note"]
    assert exported["amount_scope_note"] == project["amount_scope_note"] and exported["announcement_count"] == 1
    report = Path(result["outputs"]["report_md"]).read_text(encoding="utf-8")
    assert "实际归并项目 0 个" in report and "仅取得单份公告 1 个" in report
    assert "项目级金额未确认" in report and "当前阶段待核实" in report
    assert "多份公告已归并" not in report
    workbook = load_workbook(result["outputs"]["xlsx"], read_only=True)
    try:
        rows = list(workbook["项目清单"].values)
        assert rows[1][rows[0].index("金额")] == "项目级金额未确认"
        assert rows[1][rows[0].index("金额单位")] == "万元"
        assert rows[1][rows[0].index("当前阶段")] == "待核实"
        assert project["stage_note"] in rows[1] and project["amount_scope_note"] in rows[1]
    finally:
        workbook.close()
