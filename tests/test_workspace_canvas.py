"""画布身份与结果追问使用合成文件、隔离库和模型替身。"""
import asyncio
import hashlib
import json

import pytest

from src.api.routes import semantic_workspace as routes
from src.conversation_steering.models import ContextDelta, RawUserTurn, SteeringRequest
from src.conversation_steering.repository import SqliteSteeringRepository
from src.conversation_steering.service import ConversationSteering
from tests.database_migration_helpers import migrated_webui_database
from src.account_execution import ExecutionAuthorization, execution_context
from tests.test_workspace_conversation_stream import conversation


@pytest.fixture(autouse=True)
def execution_owner():
    with execution_context(ExecutionAuthorization("user-a", 0)):
        yield


def selected_context():
    return {
        "revision": 1, "output_id": "output-one", "representation_sha256": "a" * 64,
        "item_ref": "item_" + "b" * 64, "label": "选中的结果",
        "source_refs": [{"artifact_id": "source-one", "source_sha256": "c" * 64, "table_ref": "artifact://source-one/table/0", "row_number": 31}],
        "content": "真实选择的第31行内容",
    }


def test_selected_context_is_persisted_before_rewrite_and_idempotent(tmp_path):
    repository = SqliteSteeringRepository(str(migrated_webui_database(tmp_path / "canvas.db")))
    calls = []

    class Rewrite:
        async def rewrite(self, turn, request):
            calls.append(turn.result_context)
            assert repository.get_turn(turn.owner_id, turn.turn_id).result_context == turn.result_context
            return ContextDelta(delta_id="delta-one", owner_id=turn.owner_id, task_id=turn.task_id, inherited_revision=1, source_turn_ids=(turn.turn_id,), intent="status_question", confidence="high", normalized_text="解释结果", direct_answer="依据所选结果回答")

    request = SteeringRequest(owner_id="owner", task_id="task", revision=1, text="解释这项", idempotency_key="same", current_status="completed", result_context=selected_context())
    service = ConversationSteering(repository, Rewrite())
    result = asyncio.run(service.handle_turn(request))
    assert result.result_context is not None
    public = routes._public_steering_message(result)
    assert public["result_context"]["item_ref"] == selected_context()["item_ref"]
    assert "content" not in public["result_context"]
    assert "content" not in result.result_context.model_dump()
    assert asyncio.run(service.handle_turn(request)) == result
    changed = {**selected_context(), "item_ref": "item_" + "d" * 64}
    with pytest.raises(ValueError):
        asyncio.run(service.handle_turn(request.model_copy(update={"result_context": request.result_context.model_copy(update=changed)})))
    assert len(calls) == 1


def test_preview_item_identity_survives_search_sort_and_page(tmp_path):
    import pyarrow as pa
    import pyarrow.parquet as pq

    path = tmp_path / "result.parquet"
    pq.write_table(pa.Table.from_pylist([{"name": "重复", "value": i, "item_ref": "业务字段", "__mg_output_record_id": hashlib.sha256(str(i).encode()).hexdigest()} for i in range(40)]), path)
    args = dict(lineage_path=None, offset=0, limit=50, search="", sort_by=None, sort_direction="asc")
    full = routes._preview_result_file(path, **args)
    page = routes._preview_result_file(path, **{**args, "offset": 2, "limit": 4, "sort_by": "value", "sort_direction": "desc"})
    assert len(full["item_refs"]) == 40
    assert page["item_refs"] == list(reversed(full["item_refs"]))[2:6]
    assert page["rows"][0]["item_ref"] == "业务字段"
    assert "__mg_output_record_id" not in json.dumps(page)


@pytest.fixture
def canvas(tmp_path, monkeypatch):
    from tests.test_semantic_workspace_api import _client, _upload, _wait

    client, owner = _client(tmp_path, monkeypatch)
    upload = _upload(tmp_path)
    with client:
        created = client.post("/api/semantic-workspace/tasks", json={"objective_text": "提取谢超群工作量并输出Excel", "upload_ids": [upload.upload_id], "output_formats": ["xlsx"], "provider": "local", "model": "fixture"})
        assert created.status_code == 202, created.text
        task_id = created.json()["task_id"]
        task = _wait(client, task_id, {"completed"})
        yield client, owner, task_id, task, upload


def test_formal_preview_source_and_answer_share_persisted_identity(canvas, monkeypatch):
    monkeypatch.setattr(routes, "platform_session_valid", lambda _request: True)
    client, _, task_id, task, upload = canvas
    preview = client.get(f"/api/semantic-workspace/tasks/{task_id}/preview").json()
    assert preview["output_id"] == task["delivery"]["outputs"][0]["output_id"]
    assert preview["representation"]["kind"] == "derived_result"
    assert preview["representation"]["lineage_available"] is True
    lineage = preview["rows"][1]["__lineage"][0]
    source = client.get(f"/api/semantic-workspace/tasks/{task_id}/sources/{upload.upload_id}/preview", params={"revision": 1, "table_ref": lineage["table_ref"], "row_number": lineage["row_number"], "limit": 2})
    assert source.status_code == 200, source.text
    assert source.json()["location_status"] == "located"
    assert source.json()["total"] == 11
    calls = []

    class Rewrite:
        async def rewrite(self, turn, request):
            calls.append(turn.result_context)
            assert str(preview["rows"][1]["工作量费用"]) in turn.result_context.content
            return ContextDelta(delta_id="selected-delta", owner_id=turn.owner_id, task_id=turn.task_id, inherited_revision=1, source_turn_ids=(turn.turn_id,), intent="status_question", confidence="high", normalized_text="解释所选结果", direct_answer="所选结果的说明")

    monkeypatch.setattr(routes, "build_context_rewriter", lambda *args, **kwargs: Rewrite())
    selection = {"revision": 1, "output_id": preview["output_id"], "representation_sha256": preview["representation"]["sha256"], "item_ref": preview["item_refs"][1]}
    body = {"text": "解释此项", "result_context": selection}
    answer = client.post(f"/api/semantic-workspace/tasks/{task_id}/turns", json=body, headers={"Idempotency-Key": "selected-one"})
    assert answer.status_code == 200, answer.text
    assert answer.json()["result_context"]["source_refs"][0]["source_sha256"] == upload.sha256
    assert "content" not in answer.json()["result_context"]
    replay = client.get(f"/api/semantic-workspace/tasks/{task_id}/turns").json()
    assert replay["results"][0]["result_context"] == answer.json()["result_context"]
    assert "content" not in replay["turns"][0]["result_context"]
    assert client.get(f"/api/semantic-workspace/tasks/{task_id}").json()["messages"][0]["result_context"] == answer.json()["result_context"]
    changed = {**body, "result_context": {**selection, "item_ref": preview["item_refs"][2]}}
    assert client.post(f"/api/semantic-workspace/tasks/{task_id}/turns", json=changed, headers={"Idempotency-Key": "selected-one"}).status_code == 409
    assert client.post(f"/api/semantic-workspace/tasks/{task_id}/turns", json={"text": "解释此项"}, headers={"Idempotency-Key": "selected-one"}).status_code == 409
    stream = client.get(f"/api/semantic-workspace/tasks/{task_id}/stream?revision=1").text
    assert stream.count("event: message") == 1
    assert selection["item_ref"] in stream
    assert len(calls) == 1


def test_historical_bundle_reads_frozen_sources_and_checks_integrity(canvas):
    import io
    import zipfile
    from src.api.auth import get_store

    client, _, task_id, task, upload = canvas
    store = get_store()
    store.create_semantic_workspace_revision("user-a", task_id, objective_text="后续版本私有目标", output_formats=["xlsx"], change_summary="后续版本")
    with store._conn() as conn:
        conn.execute("UPDATE semantic_workspace_tasks SET upload_ids_json='[]' WHERE task_id=?", (task_id,))
    response = client.get(f"/api/semantic-workspace/tasks/{task_id}/bundle", params={"revision": 1, "include_sources": True})
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        source_name = next(name for name in archive.namelist() if name.startswith("sources/"))
        assert hashlib.sha256(archive.read(source_name)).hexdigest() == upload.sha256
        trace = archive.read("trace.json").decode("utf-8")
        assert "后续版本私有目标" not in trace
        assert "artifact_paths" not in trace
        assert json.loads(trace)["task"]["revision"] == 1
    record = store.get_semantic_delivery_output("user-a", task["delivery"]["outputs"][0]["output_id"])
    from pathlib import Path
    Path(record["file_path"]).write_bytes(b"changed")
    assert client.get(f"/api/semantic-workspace/tasks/{task_id}/bundle", params={"revision": 1}).status_code == 409
    assert client.get(f"/api/semantic-workspace/tasks/{task_id}/preview", params={"revision": 1}).status_code == 409


def test_csv_multiline_keeps_executor_record_number_and_original_value(tmp_path):
    from src.semantic_harness.table_executor import _read_text_rows
    path = tmp_path / "multiline.csv"
    path.write_bytes('name,note\nfirst,"first line\nsecond line"\nsecond,ok\n'.encode("utf-8"))
    assert _read_text_rows(path, delimiter=",", header_row=1) == [(2, ["first", "first line\nsecond line"]), (3, ["second", "ok"])]


def test_selected_broker_payload_and_concurrency_keep_one_frozen_call(conversation):
    from src.conversation_steering.models import FrozenResultContext, SteeringResult
    from src.conversation_steering.rewriter import BrokerContextRewriter

    database, _, _, calls, _, _, request = conversation
    request = request.model_copy(update={"result_context": FrozenResultContext(**selected_context())})

    async def scenario():
        service = ConversationSteering(SqliteSteeringRepository(database), BrokerContextRewriter())
        results = await asyncio.gather(service.handle_turn(request), service.handle_turn(request), return_exceptions=True)
        assert sum(isinstance(result, SteeringResult) for result in results) == 1
        assert len(calls) == 1
        assert selected_context()["content"] in json.dumps(calls[0], ensure_ascii=False)
        assert calls[0]["model"] == request.model
        rebuilt = ConversationSteering(SqliteSteeringRepository(database), BrokerContextRewriter())
        result = await rebuilt.handle_turn(request)
        assert result.result_context.item_ref == selected_context()["item_ref"]
        assert len(calls) == 1
    asyncio.run(scenario())


def test_selected_broker_rechecks_revision_before_network(conversation):
    from fastapi import HTTPException
    from src.conversation_steering.models import FrozenResultContext
    from src.conversation_steering.rewriter import BrokerContextRewriter

    database, store, _, calls, _, _, request = conversation
    request = request.model_copy(update={"result_context": FrozenResultContext(**selected_context())})
    store.create_semantic_workspace_revision("user-a", "task-1", objective_text="最新要求", output_formats=["json"], change_summary="变化")
    callback = lambda: routes._require_active_result("user-a", "task-1", request.result_context)
    service = ConversationSteering(SqliteSteeringRepository(database), BrokerContextRewriter(before_call=callback), before_result_call=callback)
    with pytest.raises(HTTPException) as error:
        asyncio.run(service.handle_turn(request))
    assert error.value.status_code == 409
    assert calls == []
    import sqlite3
    with sqlite3.connect(database) as conn:
        assert conn.execute("SELECT result_context_claimed FROM conversation_raw_turns").fetchone()[0] == 0


def test_result_selection_rejects_foreign_stale_or_forged_identity(canvas, monkeypatch):
    from src.api.auth import get_store

    client, owner, task_id, _, upload = canvas
    preview = client.get(f"/api/semantic-workspace/tasks/{task_id}/preview").json()
    selection = {"revision": 1, "output_id": preview["output_id"], "representation_sha256": preview["representation"]["sha256"], "item_ref": preview["item_refs"][0]}
    monkeypatch.setattr(routes, "build_context_rewriter", lambda *args, **kwargs: pytest.fail("非法选择不得进入模型"))
    for field, value, code in (("output_id", "foreign-output", 404), ("representation_sha256", "0" * 64, 409), ("item_ref", "item_" + "0" * 64, 409), ("revision", 2, 409), ("source_refs", [], 422)):
        response = client.post(f"/api/semantic-workspace/tasks/{task_id}/turns", json={"text": "解释", "result_context": {**selection, field: value}})
        assert response.status_code == code, response.text
    get_store().create_semantic_workspace_revision("user-a", task_id, objective_text="新目标", output_formats=["xlsx"], change_summary="新目标")
    assert client.post(f"/api/semantic-workspace/tasks/{task_id}/turns", json={"text": "解释", "result_context": selection}).status_code == 409
    owner["value"] = "user-b"
    assert client.get(f"/api/semantic-workspace/tasks/{task_id}/preview", params={"revision": 1}).status_code == 404
    assert client.get(f"/api/semantic-workspace/tasks/{task_id}/sources/{upload.upload_id}/preview", params={"revision": 1}).status_code == 404


def _source_preview_task(*uploads):
    """格式测试同样冻结真实来源，不能仅靠预览替身冒充任务成员。"""
    from src.api.auth import get_store

    task_id = "preview_" + uploads[0].upload_id
    get_store().create_semantic_workspace_task(
        "user-a", task_id=task_id, title="来源预览", objective_text="检查合成资料",
        upload_ids=[item.upload_id for item in uploads],
        source_refs=[{"upload_id": item.upload_id, "sha256": item.sha256} for item in uploads],
        output_formats=["json"], provider="local", model="fixture", external_api_confirmed=False,
    )
    return task_id


def test_source_full_window_multisheet_physical_rows_and_fail_closed(canvas, tmp_path, monkeypatch):
    from openpyxl import Workbook
    from src.services.upload_store import UploadStore
    from src.semantic_harness.inspectors.tabular import inspect_tabular_path
    from pathlib import Path

    client, _, task_id, _, _ = canvas
    workbook = Workbook()
    workbook.active.append(["首页"])
    sheet = workbook.create_sheet("实际数据")
    sheet.append([None])
    sheet.append(["姓名", "金额"])
    for index in range(40):
        sheet.append(["重复值", index])
    file = tmp_path / "sheets.xlsx"
    workbook.save(file)
    upload = UploadStore(root=str(tmp_path / "uploads"), max_bytes=10_000_000).save_bytes("user-a", file.name, file.read_bytes())
    report = inspect_tabular_path(artifact_id=upload.upload_id, artifact_sha256=upload.sha256, path=Path(upload.storage_path), original_name=file.name, declared_media_type=upload.media_type).model_dump(mode="json")
    task_id = _source_preview_task(upload)
    monkeypatch.setattr(routes, "_frozen_canvas_sources", lambda *args: {upload.upload_id: {"sha256": upload.sha256, "report": report}})
    url = f"/api/semantic-workspace/tasks/{task_id}/sources/{upload.upload_id}/preview"
    params = {"revision": 1, "table_ref": report["tables"][1]["table_ref"], "limit": 5, "row_number": 37, "sort_by": "金额", "sort_direction": "desc"}
    response = client.get(url, params=params)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["total"] == 40
    assert data["tables"][1]["header_row"] == 2
    assert data["location_status"] == "located"
    assert any(row["row_number"] == 37 and row["values"]["金额"] == 34 for row in data["rows"])
    assert client.get(url, params={**params, "search": "no-match"}).json()["location_status"] == "not_found"
    assert client.get(url, params={**params, "extractor_version": "old-version"}).json()["location_status"] == "version_mismatch"
    monkeypatch.setattr(routes.settings, "data_prep_preview_max_bytes", 1)
    assert client.get(url, params=params).status_code == 413


def test_source_document_version_page_and_membership_are_explicit(canvas, tmp_path, monkeypatch):
    from docx import Document
    from pypdf import PdfWriter
    from src.services.upload_store import UploadStore

    client, _, task_id, _, _ = canvas
    uploads = UploadStore(root=str(tmp_path / "uploads"), max_bytes=10_000_000)
    doc = Document()
    for index in range(40):
        doc.add_paragraph(f"合成段落{index}")
    file = tmp_path / "document.docx"
    doc.save(file)
    upload = uploads.save_bytes("user-a", file.name, file.read_bytes())
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=100, height=100)
    pdf = tmp_path / "pages.pdf"
    writer.write(pdf)
    pdf_upload = uploads.save_bytes("user-a", pdf.name, pdf.read_bytes())
    task_id = _source_preview_task(upload, pdf_upload)
    monkeypatch.setattr(routes, "_frozen_canvas_sources", lambda *args: {item.upload_id: {"sha256": item.sha256} for item in (upload, pdf_upload)})
    url = f"/api/semantic-workspace/tasks/{task_id}/sources/{upload.upload_id}/preview"
    full = client.get(url, params={"revision": 1}).json()
    assert full["total"] == 40
    target = full["elements"][34]
    assert "metadata" not in target and "raw_result_ref" not in target
    located = client.get(url, params={"revision": 1, "element_id": target["element_id"], "extractor_version": target["extractor_version"], "limit": 5}).json()
    assert located["location_status"] == "located"
    assert located["offset"] == 30
    assert client.get(url, params={"revision": 1, "element_id": target["element_id"], "extractor_version": "old"}).json()["location_status"] == "version_mismatch"
    pdf_url = f"/api/semantic-workspace/tasks/{task_id}/sources/{pdf_upload.upload_id}/preview"
    pdf_view = client.get(pdf_url, params={"revision": 1, "page": 3}).json()
    assert pdf_view["page_count"] == 3 and pdf_view["elements"] == []
    assert pdf_view["location_status"] == "located"
    assert client.get(pdf_url, params={"revision": 1, "page": 4}).json()["location_status"] == "not_found"
    assert client.get(f"/api/semantic-workspace/tasks/{task_id}/sources/not-member/preview", params={"revision": 1}).status_code == 404


def test_csv_source_window_and_executor_preserve_identical_multiline_records(canvas, tmp_path, monkeypatch):
    from src.services.upload_store import UploadStore
    from src.semantic_harness.table_executor import _read_text_rows
    from pathlib import Path

    client, _, task_id, _, _ = canvas
    raw = 'name,note\nfirst,"line one\nline two"\nsecond,ok\n'.encode("utf-8")
    upload = UploadStore(root=str(tmp_path / "uploads"), max_bytes=1_000_000).save_bytes("user-a", "lines.csv", raw)
    task_id = _source_preview_task(upload)
    monkeypatch.setattr(routes, "_frozen_canvas_sources", lambda *args: {upload.upload_id: {"sha256": upload.sha256}})
    response = client.get(f"/api/semantic-workspace/tasks/{task_id}/sources/{upload.upload_id}/preview", params={"revision": 1})
    assert response.status_code == 200, response.text
    rows = response.json()["rows"]
    assert [(row["row_number"], list(row["values"].values())) for row in rows] == _read_text_rows(Path(upload.storage_path), delimiter=",", header_row=1)
    assert rows[0]["values"]["note"] == "line one\nline two"


def test_image_source_preview_is_original_without_ocr(canvas, tmp_path, monkeypatch):
    import io
    from PIL import Image
    from src.services.upload_store import UploadStore
    from src.parsers.image import ImageParser

    client, _, task_id, _, _ = canvas
    buffer = io.BytesIO()
    picture = Image.new("RGB", (12, 8), "white")
    exif = Image.Exif()
    exif[274] = 6
    picture.save(buffer, format="JPEG", exif=exif)
    payload = buffer.getvalue()
    uploads = UploadStore(root=str(tmp_path / "uploads"), max_bytes=4096)
    upload = uploads.save_bytes("user-a", "rotated.jpg", payload, verify_magic=True)
    task_id = _source_preview_task(upload)
    monkeypatch.setattr(routes, "_frozen_canvas_sources", lambda *args: {upload.upload_id: {"sha256": upload.sha256}})
    monkeypatch.setattr(ImageParser, "parse", lambda *args: pytest.fail("原件预览禁止触发 OCR"))
    url = f"/api/semantic-workspace/tasks/{task_id}/sources/{upload.upload_id}/preview"
    response = client.get(url, params={"revision": 1})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["kind"] == "image"
    assert (result["image_width"], result["image_height"], result["image_orientation"]) == (12, 8, 6)
    assert result["representation"]["kind"] == "source"
    assert result["sha256"] == hashlib.sha256(payload).hexdigest()
    assert result["elements"] == [] and result["page_count"] == 1
    assert client.get(url, params={"revision": 1, "page": 2}).json()["location_status"] == "not_found"
    assert client.get(url, params={"revision": 1, "element_id": "unverified"}).json()["location_status"] == "not_found"


def test_published_representation_does_not_follow_latest_attempt(canvas, monkeypatch):
    from src.api.auth import get_store

    client, _, task_id, _, _ = canvas
    first = client.get(f"/api/semantic-workspace/tasks/{task_id}/preview").json()
    monkeypatch.setattr(get_store(), "latest_semantic_harness_artifact_paths", lambda *args: pytest.fail("正式发布不能选择最新Attempt"))
    again = client.get(f"/api/semantic-workspace/tasks/{task_id}/preview")
    assert again.status_code == 200, again.text
    assert again.json() == first


def test_reference_whitelist_does_not_create_or_strip_invalid_evidence(monkeypatch):
    monkeypatch.setattr(routes, "_canvas_upload", lambda *args: None)
    from fastapi import HTTPException
    assert routes._canvas_source_refs("owner", [], {}) == []
    with pytest.raises(HTTPException):
        routes._canvas_source_refs("owner", [{"artifact_id": "missing", "page": 2}], {})
    public = routes._canvas_source_refs("owner", [{"artifact_id": "source", "element_id": "el-one", "page": 2, "quote": "引用内容", "raw_result_ref": "private", "metadata": {"hidden": True}, "location": {"kind": "docx_paragraph", "paragraph": 2, "hidden": True}}], {"source": {"sha256": "a" * 64}})
    assert public == [{"artifact_id": "source", "source_sha256": "a" * 64, "element_id": "el-one", "page": 2, "location": {"kind": "docx_paragraph", "paragraph": 2}}]


def test_unknown_selected_turn_survives_rebuild_without_resending(conversation, monkeypatch):
    from src.conversation_steering.models import FrozenResultContext
    from src.conversation_steering.rewriter import BrokerContextRewriter

    database, _, broker, calls, _, _, request = conversation
    request = request.model_copy(update={"result_context": FrozenResultContext(**selected_context())})
    async def crash(**kwargs):
        raise asyncio.CancelledError()
    monkeypatch.setattr(broker, "relay", crash)
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(ConversationSteering(SqliteSteeringRepository(database), BrokerContextRewriter()).handle_turn(request))
    repository = SqliteSteeringRepository(database)
    assert repository.list_turns("user-a", "task-1")[0].result_context == request.result_context
    with pytest.raises(ValueError, match="禁止自动重复"):
        asyncio.run(ConversationSteering(repository, BrokerContextRewriter()).handle_turn(request))
    assert calls == []
    assert broker.list_usage("user-a", task_id="task-1", revision=1)[0]["status"] == "unknown"


def test_source_formula_without_cache_and_tampered_source_are_rejected(canvas, tmp_path, monkeypatch):
    from openpyxl import Workbook
    from src.services.upload_store import UploadStore
    from pathlib import Path

    client, _, task_id, _, _ = canvas
    workbook = Workbook()
    workbook.active.append(["value"])
    workbook.active.append(["=1+1"])
    file = tmp_path / "formula.xlsx"
    workbook.save(file)
    upload = UploadStore(root=str(tmp_path / "uploads"), max_bytes=1_000_000).save_bytes("user-a", file.name, file.read_bytes())
    task_id = _source_preview_task(upload)
    monkeypatch.setattr(routes, "_frozen_canvas_sources", lambda *args: {upload.upload_id: {"sha256": upload.sha256}})
    url = f"/api/semantic-workspace/tasks/{task_id}/sources/{upload.upload_id}/preview?revision=1"
    assert client.get(url).status_code == 409
    Path(upload.storage_path).write_bytes(b"modified")
    assert client.get(url).status_code == 409


def test_bundle_failure_removes_only_its_temporary_file(canvas, monkeypatch):
    import zipfile
    from pathlib import Path

    client, _, task_id, _, _ = canvas
    root = Path(routes.settings.semantic_execution_root) / "_bundles"
    root.mkdir(exist_ok=True)
    existing = root / "keep.txt"
    existing.write_text("保留既有文件", encoding="utf-8")
    original = zipfile.ZipFile.open
    def fail(self, name, mode="r", *args, **kwargs):
        if mode == "w":
            raise OSError("synthetic archive failure")
        return original(self, name, mode, *args, **kwargs)
    monkeypatch.setattr(zipfile.ZipFile, "open", fail)
    with pytest.raises(OSError, match="synthetic archive failure"):
        client.get(f"/api/semantic-workspace/tasks/{task_id}/bundle")
    assert list(root.iterdir()) == [existing]


def test_legacy_unpublished_preview_keeps_verified_lineage_without_selection(canvas, monkeypatch):
    from src.api.auth import get_store

    client, _, task_id, _, _ = canvas
    monkeypatch.setattr(get_store(), "latest_semantic_delivery", lambda *args: None)
    response = client.get(f"/api/semantic-workspace/tasks/{task_id}/preview")
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["output_id"] is None
    assert preview["representation"]["kind"] == "intermediate"
    assert preview["rows"][0]["__lineage"]
    assert preview["item_refs"] == [None] * len(preview["rows"])


def test_bundle_rejects_incomplete_manifest_and_filters_admin_trace(canvas, monkeypatch):
    from src.api.auth import get_store
    import io
    import zipfile

    client, _, task_id, _, _ = canvas
    store = get_store()
    store.append_semantic_workspace_event("user-a", task_id, stage="execute", event_type="admin_note", summary="仅限管理员的诊断", details={"audience": "admin", "refs": {"prompt": "内部提示正文"}})
    store.append_semantic_workspace_event("user-a", task_id, stage="execute", event_type="visible_note", summary="用户可见的正常进度", details={"audience": "all", "refs": {"prompt": "不该公开的提示"}, "action": {"command": "internal", "parameters": {"private": "内部执行参数"}}})
    response = client.get(f"/api/semantic-workspace/tasks/{task_id}/bundle", params={"include_sources": True})
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        trace = archive.read("trace.json").decode("utf-8")
    assert "用户可见的正常进度" in trace
    assert all(value not in trace for value in ("仅限管理员", "内部提示正文", "不该公开的提示", "内部执行参数"))
    original = store.latest_semantic_delivery
    monkeypatch.setattr(store, "latest_semantic_delivery", lambda *args: {**original(*args), "source_artifact_hashes": {}})
    assert client.get(f"/api/semantic-workspace/tasks/{task_id}/bundle", params={"include_sources": True}).status_code == 409


def test_non_broker_selected_turn_claim_is_atomic_across_threads(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    import threading
    from src.conversation_steering.models import SteeringResult

    database = str(migrated_webui_database(tmp_path / "concurrent.db"))
    barrier = threading.Barrier(2)
    calls = []
    class Rewrite:
        async def rewrite(self, turn, request):
            calls.append(turn.turn_id)
            return ContextDelta(delta_id="local-delta", owner_id=turn.owner_id, task_id=turn.task_id, inherited_revision=1, source_turn_ids=(turn.turn_id,), intent="status_question", confidence="high", normalized_text="解释", direct_answer="本地解释")
    request = SteeringRequest(owner_id="owner", task_id="task", revision=1, text="解释", idempotency_key="once", current_status="completed", result_context=selected_context())
    repository = SqliteSteeringRepository(database)
    repository.save_turn(RawUserTurn(turn_id="fixed-turn", owner_id="owner", task_id="task", revision=1, text="解释", idempotency_key="once", result_context=selected_context()))
    def run():
        try:
            return asyncio.run(ConversationSteering(SqliteSteeringRepository(database), Rewrite(), before_result_call=lambda: barrier.wait(timeout=10)).handle_turn(request))
        except ValueError as exc:
            return exc
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: run(), range(2)))
    assert sum(isinstance(result, SteeringResult) for result in results) == 1
    assert len(calls) == 1
    assert asyncio.run(ConversationSteering(SqliteSteeringRepository(database), Rewrite()).handle_turn(request)).answer == "本地解释"
    assert len(calls) == 1


def test_non_broker_unknown_selected_turn_does_not_retry_after_rebuild(tmp_path):
    database = str(migrated_webui_database(tmp_path / "unknown.db"))
    calls = []
    class Rewrite:
        async def rewrite(self, turn, request):
            calls.append(turn.turn_id)
            raise TimeoutError("模型响应未知")
    request = SteeringRequest(owner_id="owner", task_id="task", revision=1, text="解释", idempotency_key="once", current_status="completed", result_context=selected_context())
    with pytest.raises(TimeoutError):
        asyncio.run(ConversationSteering(SqliteSteeringRepository(database), Rewrite()).handle_turn(request))
    with pytest.raises(ValueError, match="禁止自动重复"):
        asyncio.run(ConversationSteering(SqliteSteeringRepository(database), Rewrite()).handle_turn(request))
    assert len(calls) == 1


@pytest.mark.parametrize("failure", ["timeout", "server_error"])
def test_selected_instructor_http_is_not_retried_after_unknown(tmp_path, monkeypatch, failure):
    from types import SimpleNamespace
    import httpx
    from src.conversation_steering import rewriter

    database = str(migrated_webui_database(tmp_path / "sdk-unknown.db"))
    sends = []
    def receive(request):
        sends.append(json.loads(request.content))
        if failure == "timeout":
            raise httpx.ReadTimeout("合成请求已收到，响应未知", request=request)
        return httpx.Response(503, json={"error": {"message": "合成服务错误"}})
    original_client = httpx.AsyncClient
    class IsolatedClient(original_client):
        def __init__(self, **kwargs):
            super().__init__(**kwargs, transport=httpx.MockTransport(receive))
    monkeypatch.setattr(rewriter.httpx, "AsyncClient", IsolatedClient)
    async def authorized_checkpoint(_message):
        return None
    monkeypatch.setattr(rewriter, "execution_http_checkpoint_async", authorized_checkpoint)
    connection = SimpleNamespace(timeout=5, trust_env=False, api_key="fixture", base_url="https://fixture.invalid/v1", provider="local", model="fixture", extra_body={})
    monkeypatch.setattr(rewriter, "get_provider", lambda: SimpleNamespace(resolve_model=lambda *args, **kwargs: connection))
    request = SteeringRequest(owner_id="user-a", task_id="task", revision=1, text="解释", idempotency_key="once", current_status="completed", result_context=selected_context())
    def service():
        return ConversationSteering(SqliteSteeringRepository(database), rewriter.InstructorContextRewriter(provider="local", model="fixture"))
    with pytest.raises(Exception) as failure_result:
        asyncio.run(service().handle_turn(request))
    assert len(sends) == 1, type(failure_result.value).__name__
    assert selected_context()["content"] in sends[0]["messages"][-1]["content"]
    with pytest.raises(ValueError, match="禁止自动重复"):
        asyncio.run(service().handle_turn(request))
    assert len(sends) == 1


def test_public_preview_filters_internal_refs_but_internal_selection_keeps_them(canvas, tmp_path, monkeypatch):
    client, _, task_id, _, _ = canvas
    path = tmp_path / "document.json"
    ref = {"artifact_id": "source", "element_id": "element", "page": 1, "quote": "用户证据正文", "raw_result_ref": "private-reference", "metadata": {"prompt": "内部字段"}, "location": {"kind": "docx_paragraph", "paragraph": 1, "private": "内部字段"}}
    path.write_text(json.dumps({"passages": [{"passage_id": "passage", "label": "条款", "text": "合法正文", "evidence_refs": [ref]}]}, ensure_ascii=False), encoding="utf-8")
    identity = {"task_id": task_id, "revision": 1, "run_id": "run", "delivery_id": "delivery", "output_id": "output", "representation": {"kind": "output", "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}}
    monkeypatch.setattr(routes, "_canvas_result", lambda *args: (path, None, {}, identity))
    response = client.get(f"/api/semantic-workspace/tasks/{task_id}/preview")
    assert response.status_code == 200, response.text
    public = response.json()["items"][0]["evidence_refs"][0]
    assert public["quote"] == "用户证据正文"
    assert "raw_result_ref" not in public and "metadata" not in public and "private" not in public["location"]
    raw = routes._preview_result_file(path, lineage_path=None, offset=0, limit=100, search="", sort_by=None, sort_direction="asc")
    assert raw["items"][0]["evidence_refs"][0] == ref


def test_selected_compressed_result_obeys_preview_expansion_limit(canvas, tmp_path, monkeypatch):
    import zipfile

    client, _, task_id, _, _ = canvas
    path = tmp_path / "compressed.xlsx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "x" * 20_000)
    monkeypatch.setattr(routes.settings, "data_prep_preview_max_bytes", 10_000)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setattr(routes, "_canvas_result", lambda *args: (path, None, {}, {"representation": {"sha256": digest}}))
    monkeypatch.setattr(routes, "_preview_result_file", lambda *args, **kwargs: pytest.fail("超界压缩结果不得解析"))
    monkeypatch.setattr(routes, "build_context_rewriter", lambda *args, **kwargs: pytest.fail("超界压缩结果不得调用模型"))
    selection = {"revision": 1, "output_id": "output", "representation_sha256": digest, "item_ref": "item_" + "b" * 64}
    response = client.post(f"/api/semantic-workspace/tasks/{task_id}/turns", headers={"Idempotency-Key": "compressed"}, json={"text": "解释", "result_context": selection})
    assert response.status_code == 413, response.text
