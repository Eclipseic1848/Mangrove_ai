"""合成任务走真实初稿接受和 Publisher，不调用模型。"""
import pytest
from types import SimpleNamespace

from src.account_execution import ExecutionAuthorization, execution_context
from src.agentic_runtime.draft_snapshot import freeze_draft
from src.agentic_runtime.models import RuntimeTaskConfig, RuntimeVersion, RuntimeStatus, PermissionProfile
from src.agentic_runtime.repository import AgenticRuntimeRepository
from src.api.store import WebUIStore
from src.api.workspace_draft_acceptance import accept_draft
from src.config.settings import settings
from src.services.managed_paths import ManagedPathCodec
from tests.account_execution_helpers import seed_execution_owner
from tests.database_migration_helpers import migrated_webui_database


def test_draft_review_requires_explicit_same_run_decision_and_survives_restart(draft_task, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from src.api.auth import get_current_user
    from src.api.routes import semantic_workspace as routes
    from src.api.semantic_workspace_runtime import SemanticWorkspaceManager
    store, root, draft = draft_task
    repo = AgenticRuntimeRepository(store.db_path)
    (root / "session").mkdir()
    session = root / "session" / "run.jsonl"
    session.write_text('{}\n', encoding="utf-8")
    repo.update("a", "t", 1, status=RuntimeStatus.NEEDS_INPUT, session_file="session/run.jsonl")
    repo.append_event("a", "t", 1, event_type="draft.ready", summary="等待决定", details={"draft_id": draft["draft_id"]})
    store.publish_workspace_question("a", "t", {
        "kind": "plan", "purpose": "control", "continuation": "unavailable",
        "question_id": "draft-review:" + draft["draft_id"], "draft_review_id": draft["draft_id"],
        "draft_review_run_id": "r", "prompt": "初稿已生成", "options": [], "allow_free_text": False,
    }, expected_revision=1)
    request = SimpleNamespace(user_id="a", task_id="t", revision=1)
    assert SemanticWorkspaceManager._draft_review_required(request)
    assert store.list_pending_semantic_workspace_tasks() == []
    app = FastAPI()
    app.include_router(routes.router)
    owner = {"user_id": "a", "execution_generation": 0}
    app.dependency_overrides[get_current_user] = lambda: owner
    enqueued = []
    monkeypatch.setattr(routes, "get_semantic_workspace_manager", lambda: SimpleNamespace(enqueue=lambda *args: enqueued.append(args)))
    payload = {"expected_revision": 1, "draft_id": draft["draft_id"]}
    with TestClient(app) as client:
        assert client.get("/api/semantic-workspace/tasks/t/draft?revision=1").json()["draft"]["review_waiting"]
        owner["user_id"] = "b"
        assert client.post("/api/semantic-workspace/tasks/t/draft/verify", json=payload).status_code == 404
        owner["user_id"] = "a"
        repo.update("a", "t", 1, session_file="missing.jsonl")
        assert client.post("/api/semantic-workspace/tasks/t/draft/verify", json=payload).status_code == 409
        assert enqueued == []
        repo.update("a", "t", 1, session_file="session/run.jsonl")
        response = client.post("/api/semantic-workspace/tasks/t/draft/verify", json=payload)
        assert response.status_code == 200, response.text
        assert store.get_semantic_workspace_task("a", "t")["active_revision"] == 1
        assert not SemanticWorkspaceManager._draft_review_required(request)
        store.update_semantic_workspace_task("a", "t", status="running")
        assert client.post("/api/semantic-workspace/tasks/t/draft/verify", json=payload).status_code == 200
        assert len(enqueued) == 1
        repo.update("a", "t", 1, run_id="other-run")
        assert SemanticWorkspaceManager._draft_review_required(request)
        assert client.post("/api/semantic-workspace/tasks/t/draft/verify", json=payload).status_code in {404, 409}
        repo.update("a", "t", 1, run_id="r")
        store.update_semantic_workspace_task("a", "t", status="cancelled", cancel_requested=True)
        assert client.post("/api/semantic-workspace/tasks/t/draft/verify", json=payload).status_code == 409
        assert len(enqueued) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("initial_status", ["failed", "cancelled"])
@pytest.mark.parametrize("name,content,fmt", [("document.json", '{"name":"测试"}', "json"),
                                           ("table.csv", "name,value\na,1\n", "csv"),
                                           ("collection.md", "# 合成来源分析", "markdown")])
async def test_accept_draft_publishes_same_bytes_without_rerunning_and_replays(tmp_path, monkeypatch, name, content, fmt, initial_status):
    db = migrated_webui_database(tmp_path / "db.sqlite")
    seed_execution_owner(db, "a")
    monkeypatch.setattr(settings, "webui_db_path", str(db))
    monkeypatch.setattr(settings, "semantic_execution_root", str(tmp_path))
    store = WebUIStore(str(db), semantic_paths=ManagedPathCodec(tmp_path / "deliveries", legacy_anchor=("data", "semantic-executions")))
    # 执行安全点复用认证 Store；绑定本例临时库，避免借用其他测试留下的全局实例。
    import src.api.auth as auth
    monkeypatch.setattr(auth, "_store", store)
    root = tmp_path / "run"
    (root / "output").mkdir(parents=True)
    (root / "output" / name).write_text(content, encoding="utf-8")
    with execution_context(ExecutionAuthorization("a", 0)):
        store.create_semantic_workspace_task("a", task_id="t", title="合成任务", objective_text="整理来源",
                                            upload_ids=[], output_formats=[fmt], provider="local", model=None,
                                            external_api_confirmed=False)
        store.update_semantic_workspace_task("a", "t", status=initial_status, cancel_requested=initial_status == "cancelled")
        repository = AgenticRuntimeRepository(db)
        repository.register(RuntimeTaskConfig(user_id="a", task_id="t", revision=1,
                                              runtime_version=RuntimeVersion.PI, permission_profile=PermissionProfile.STANDARD))
        repository.update("a", "t", 1, status=RuntimeStatus.FAILED, run_id="r", workspace_root=root)
        draft = freeze_draft(root, owner_id="a", task_id="t", revision=1, run_id="r", formats=(fmt,))
        result = await accept_draft(store=store, manager=None, output_root=tmp_path / "deliveries",
                                    owner_id="a", task_id="t", source_revision=1, draft_id=draft["draft_id"])
        assert result["revision"] == 2
        delivery = store.get_semantic_delivery("a", result["delivery_id"])
        assert delivery["provenance"]["verification_status"] == "inconclusive"
        assert delivery["outputs"][0]["sha256"] == draft["files"][0]["sha256"]
        replay = await accept_draft(store=store, manager=None, output_root=tmp_path / "deliveries",
                                   owner_id="a", task_id="t", source_revision=1, draft_id=draft["draft_id"])
        assert replay == result
        assert repository.get("a", "t", 1)["status"] == RuntimeStatus.FAILED


@pytest.fixture
def draft_task(tmp_path, monkeypatch, request):
    db = migrated_webui_database(tmp_path / "db.sqlite")
    seed_execution_owner(db, "a")
    monkeypatch.setattr(settings, "webui_db_path", str(db))
    monkeypatch.setattr(settings, "semantic_execution_root", str(tmp_path))
    monkeypatch.setattr(settings, "data_prep_upload_root", str(tmp_path / "uploads"))
    store = WebUIStore(str(db), semantic_paths=ManagedPathCodec(tmp_path, legacy_anchor=("data", "semantic-executions")))
    import src.api.auth as auth
    monkeypatch.setattr(auth, "_store", store)
    root = tmp_path / "run"
    (root / "output").mkdir(parents=True)
    (root / "output" / "draft.json").write_text('{"name":"合成初稿"}', encoding="utf-8")
    with execution_context(ExecutionAuthorization("a", 0)):
        from src.source_acquisition.reuse import uploads
        upload = uploads().save_bytes("a", "source.csv", b"name,value\nA,2\n", media_type="text/csv")
        store.create_semantic_workspace_task("a", task_id="t", title="合成任务", objective_text=getattr(request, 'param', {}).get('objective', '整理'),
            source_contract={'notification_user_text': getattr(request, 'param', {}).get('user_text', getattr(request, 'param', {}).get('objective', '整理'))},
            upload_ids=[upload.upload_id], source_refs=[{"upload_id": upload.upload_id, "sha256": upload.sha256}],
            output_formats=["json"], provider="local", model=None, external_api_confirmed=False)
        store.update_semantic_workspace_task("a", "t", status="failed")
        repository = AgenticRuntimeRepository(db)
        repository.register(RuntimeTaskConfig(user_id="a", task_id="t", revision=1,
            runtime_version=RuntimeVersion.PI, permission_profile=PermissionProfile.STANDARD))
        repository.update("a", "t", 1, status=RuntimeStatus.FAILED, run_id="r", workspace_root=root)
        draft = freeze_draft(root, owner_id="a", task_id="t", revision=1, run_id="r", formats=("json",))
        yield store, root, draft


@pytest.mark.asyncio
@pytest.mark.parametrize("race", ["stopped", "cancel", "complete", "deleted", "not_stopped"])
async def test_accept_stops_running_task_and_honors_competing_terminal_action(draft_task, race):
    store, root, draft = draft_task
    store.update_semantic_workspace_task("a", "t", status="running")
    class Manager:
        async def cancel(self, owner, task, *, for_revision):
            assert for_revision
            if race == "cancel":
                store.request_semantic_workspace_cancellation(owner, task)
            store.update_semantic_workspace_task(owner, task,
                status="completed" if race == "complete" else "running" if race == "not_stopped" else "cancelled",
                **({"deleted_at": "2026-09-14T00:00:00"} if race == "deleted" else {}))
    params = dict(store=store, manager=Manager(), output_root=root.parent,
                  owner_id="a", task_id="t", source_revision=1, draft_id=draft["draft_id"])
    if race == "stopped":
        result = await accept_draft(**params)
        assert result["revision"] == 2
    else:
        with pytest.raises(ValueError):
            await accept_draft(**params)
        assert store.get_semantic_workspace_task("a", "t")["active_revision"] == 1


@pytest.mark.asyncio
async def test_cancel_at_commit_point_blocks_publication_and_replay(draft_task, monkeypatch):
    from src.delivery_publishing.repository import DeliveryPublishingRepository
    store, root, draft = draft_task
    original = DeliveryPublishingRepository.begin_commit
    def cancel_then_commit(self, command, **kwargs):
        store.request_semantic_workspace_cancellation("a", "t")
        return original(self, command, **kwargs)
    monkeypatch.setattr(DeliveryPublishingRepository, "begin_commit", cancel_then_commit)
    with pytest.raises(ValueError, match="取消"):
        await accept_draft(store=store, manager=None, output_root=root.parent,
            owner_id="a", task_id="t", source_revision=1, draft_id=draft["draft_id"])
    import sqlite3
    with sqlite3.connect(store.db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM formal_delivery_runs").fetchone()[0] == 0


@pytest.mark.asyncio
async def test_stop_recovery_does_not_create_a_new_user_cancel(draft_task, monkeypatch):
    import asyncio
    from src.api.semantic_workspace_runtime import SemanticWorkspaceManager
    store, root, draft = draft_task
    manager = SemanticWorkspaceManager()
    store.update_semantic_workspace_task("a", "t", status="cancelling", cancel_requested=True)
    generation = store.get_semantic_workspace_task("a", "t")["cancel_generation"]
    sleeps = 0
    async def one_iteration(_seconds):
        nonlocal sleeps
        sleeps += 1
        if sleeps > 1:
            raise asyncio.CancelledError
    class Kernel:
        async def cancel(self, *args):
            pass
    monkeypatch.setattr(asyncio, "sleep", one_iteration)
    monkeypatch.setattr(manager, "_candidate_verification_module", lambda: None)
    monkeypatch.setattr(manager, "_kernel_for_run", lambda *args: Kernel())
    # 后台维护不继承 HTTP 身份，只能恢复已有停止意图。
    from contextvars import Context
    with pytest.raises(asyncio.CancelledError):
        await asyncio.create_task(manager._maintenance_loop(), context=Context())
    task = store.get_semantic_workspace_task("a", "t")
    assert task["status"] == "cancelled"
    assert task["cancel_generation"] == generation


@pytest.mark.asyncio
async def test_accept_with_real_stop_and_background_recovery_publishes_once(draft_task, monkeypatch):
    import asyncio
    from contextvars import Context
    from src.api import semantic_workspace_runtime as runtime_module
    store, root, draft = draft_task
    manager = runtime_module.SemanticWorkspaceManager()
    store.update_semantic_workspace_task("a", "t", status="running")
    recovery_entered = asyncio.Event()
    calls = 0
    class Kernel:
        async def cancel(self, *args):
            nonlocal calls
            calls += 1
            if calls == 1:
                await asyncio.wait_for(recovery_entered.wait(), timeout=5)
            else:
                recovery_entered.set()
    monkeypatch.setattr(runtime_module, "REVERIFICATION_RECOVERY_POLL_SECONDS", 0.01)
    monkeypatch.setattr(manager, "_candidate_verification_module", lambda: None)
    monkeypatch.setattr(manager, "_kernel_for_run", lambda *args: Kernel())
    maintenance = asyncio.create_task(manager._maintenance_loop(), context=Context())
    params = dict(store=store, manager=manager, output_root=root.parent, owner_id="a", task_id="t",
                  source_revision=1, draft_id=draft["draft_id"])
    try:
        result = await accept_draft(**params)
        assert result["status"] == "completed"
        assert await accept_draft(**params) == result
        assert store.get_semantic_workspace_task("a", "t")["active_revision"] == 2
        assert recovery_entered.is_set()
    finally:
        maintenance.cancel()
        with pytest.raises(asyncio.CancelledError):
            await maintenance


def test_draft_routes_preview_download_accept_and_owner_isolation(draft_task, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from src.api.auth import get_current_user
    from src.api.routes import semantic_workspace as routes
    from src.api.routes import semantic_deliveries
    store, root, draft = draft_task
    repo = AgenticRuntimeRepository(store.db_path)
    repo.append_event("a", "t", 1, event_type="draft.ready", summary="初稿可查看", details={"draft_id": draft["draft_id"]})
    app = FastAPI()
    app.include_router(routes.router)
    app.include_router(semantic_deliveries.router)
    owner = {"user_id": "a", "execution_generation": 0}
    app.dependency_overrides[get_current_user] = lambda: owner
    monkeypatch.setattr(routes, "get_semantic_workspace_manager", lambda: None)
    with TestClient(app) as client:
        response = client.get("/api/semantic-workspace/tasks/t/draft?revision=1")
        assert response.status_code == 200, response.text
        item = response.json()["draft"]["files"][0]
        assert "合成初稿" in item["preview"]
        download = client.get(item["download_url"])
        assert download.status_code == 200
        assert download.headers["X-Mangrove-Artifact-Status"] == "unverified-draft"
        owner["user_id"] = "b"
        assert client.get(item["download_url"]).status_code == 404
        owner["user_id"] = "a"
        payload = {"expected_revision": 1, "draft_id": draft["draft_id"], "accept_unverified": True}
        accepted = client.post("/api/semantic-workspace/tasks/t/draft/accept", json=payload)
        assert accepted.status_code == 200, accepted.text
        assert client.post("/api/semantic-workspace/tasks/t/draft/accept", json=payload).json() == accepted.json()
        delivery = store.get_semantic_delivery("a", accepted.json()["delivery_id"])
        formal = client.get(f"/api/semantic-deliveries/outputs/{delivery['outputs'][0]['output_id']}")
        assert formal.status_code == 200, formal.text
        assert formal.content == download.content
        assert client.get("/api/semantic-workspace/tasks/t/draft?revision=2").json()["draft"]["draft_id"] == draft["draft_id"]
        context = routes._public_runtime("a", "t", 2)
        assert context["runtime_version"] == "pi"
        assert context["external_api_confirmed"] is False
        assert context.get("run_id") is None
        assert routes._revision_runtime_context("a", "t", 2)["run_id"] == "r"


def test_reopening_accepted_draft_keeps_source_execution_usage(draft_task, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from src.api.auth import get_current_user
    from src.api.routes import semantic_workspace as routes
    store, _, draft = draft_task
    monkeypatch.setattr(routes, "get_semantic_workspace_manager", lambda: None)
    for event_type, details in [("runtime.preparing", {}), ("provider.usage", {"input_tokens": 120, "output_tokens": 30, "total_tokens": 150}), ("run.failed", {})]:
        store.append_semantic_workspace_event("a", "t", stage="execute", event_type=event_type, summary="合成执行记录",
            details={"revision": 1, "run_id": "r", "runtime_event_type": event_type, **details})
    app = FastAPI(); app.include_router(routes.router)
    owner = {"user_id": "a", "execution_generation": 0}
    app.dependency_overrides[get_current_user] = lambda: owner
    with TestClient(app) as client:
        before = client.get("/api/semantic-workspace/tasks/t?revision=1")
        assert before.status_code == 200
        assert before.json()["work_session"]["usage"]["total_tokens"] == 150
        accepted = client.post("/api/semantic-workspace/tasks/t/draft/accept", json={"expected_revision": 1, "draft_id": draft["draft_id"], "accept_unverified": True})
        assert accepted.status_code == 200
        for _ in range(2):
            reopened = client.get("/api/semantic-workspace/tasks/t").json()
            assert reopened["viewing_revision"] == 2
            assert reopened["work_session"] is not None
            assert reopened["work_session"]["revision"] == 1
            assert reopened["work_session"]["usage"] == before.json()["work_session"]["usage"]
            assert reopened["agentic_runtime"].get("run_id") is None
        source_contract = store.get_semantic_workspace_revision("a", "t", 2)["source_contract"]
        store.create_semantic_workspace_revision("a", "t", objective_text="合成后续任务", output_formats=["json"],
            change_summary="后续修订", source_contract=source_contract, expected_revision=3)
        later = client.get("/api/semantic-workspace/tasks/t").json()
        assert later["viewing_revision"] == 3
        assert later["work_session"] is None
        owner["user_id"] = "b"
        assert client.get("/api/semantic-workspace/tasks/t").status_code == 404


def test_csv_draft_preview_is_bounded_json_not_a_download(draft_task, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from src.api.auth import get_current_user
    from src.api.routes import semantic_workspace as routes
    store, root, _ = draft_task
    (root / "output" / "draft.json").unlink()
    (root / "output" / "table.csv").write_text('部门,说明\n华东,"交通,住宿"\n' + '华南,测试\n' * 55, encoding="utf-8")
    draft = freeze_draft(root, owner_id="a", task_id="t", revision=1, run_id="r", formats=("csv",))
    AgenticRuntimeRepository(store.db_path).append_event("a", "t", 1, event_type="draft.ready", summary="初稿", details={"draft_id": draft["draft_id"]})
    app = FastAPI()
    app.include_router(routes.router)
    app.dependency_overrides[get_current_user] = lambda: {"user_id": "a", "execution_generation": 0}
    with TestClient(app) as client:
        response = client.get("/api/semantic-workspace/tasks/t/draft?revision=1")
        assert response.status_code == 200
        assert "content-disposition" not in response.headers
        table = response.json()["draft"]["files"][0]["preview_table"]
        assert table[:2] == [["部门", "说明"], ["华东", "交通,住宿"]]
        assert len(table) == 51
        assert response.json()["draft"]["files"][0]["preview_truncated"] is True


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx"])
def test_office_draft_preview_uses_frozen_files_without_download(draft_task, fmt):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from src.api.auth import get_current_user
    from src.api.routes import semantic_workspace as routes
    store, root, _ = draft_task
    (root / "output" / "draft.json").unlink()
    path = root / "output" / f"report.{fmt}"
    if fmt == "pdf":
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=300)
        with path.open("wb") as stream:
            writer.write(stream)
    elif fmt == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph("合成文档正文")
        document.save(path)
    else:
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["部门", "金额"])
        workbook.active.append(["工程部", 8458])
        workbook.active.append(["长单元格", "合成内容" * 8000])
        workbook.save(path)
        workbook.close()
    draft = freeze_draft(root, owner_id="a", task_id="t", revision=1, run_id="r", formats=(fmt,))
    AgenticRuntimeRepository(store.db_path).append_event("a", "t", 1, event_type="draft.ready", summary="初稿", details={"draft_id": draft["draft_id"]})
    app = FastAPI()
    app.include_router(routes.router)
    owner = {"user_id": "a", "execution_generation": 0}
    app.dependency_overrides[get_current_user] = lambda: owner
    with TestClient(app) as client:
        response = client.get("/api/semantic-workspace/tasks/t/draft?revision=1")
        assert response.status_code == 200, response.text
        item = response.json()["draft"]["files"][0]
        if fmt == "pdf":
            preview = client.get(item["page_preview_url"] + "&page=1")
            assert preview.status_code == 200, preview.text
            assert preview.json()["page_count"] == 1
            assert preview.json()["image"].startswith("data:image/png;base64,")
            assert "content-disposition" not in preview.headers
            assert client.get(item["page_preview_url"] + "&page=2").status_code == 422
            owner["user_id"] = "b"
            assert client.get(item["page_preview_url"] + "&page=1").status_code == 404
        elif fmt == "docx":
            assert "合成文档正文" in item["preview"]
        else:
            assert item["preview_table"][:2] == [["部门", "金额"], ["工程部", "8458"]]
            assert sum(len(cell) for row in item["preview_table"] for cell in row) <= 16000
            assert item["preview_truncated"] is True


@pytest.mark.asyncio
async def test_no_progress_leaves_task_waiting_with_draft_not_failed_or_requeued(draft_task, monkeypatch):
    from src.api.semantic_workspace_runtime import SemanticWorkspaceManager
    from src.agentic_runtime.pi_runtime import PiVerificationStalled
    store, root, draft = draft_task
    manager = SemanticWorkspaceManager()
    async def stalled(*args):
        raise PiVerificationStalled("连续两轮验证没有进展")
    async def stopped(*args):
        return True
    monkeypatch.setattr(manager, "_run_pi_task", stalled)
    monkeypatch.setattr(manager, "_confirm_runtime_stopped", stopped)
    await manager._run_task_authorized("a", "t")
    task = store.get_semantic_workspace_task("a", "t")
    assert task["status"] == "needs_input"
    assert task["failure"]["error_code"] == "VERIFICATION_STALLED"
    assert task["active_revision"] == 1
    assert (root / "drafts" / draft["draft_id"] / "draft.json").is_file()


@pytest.mark.parametrize("state,message", [("cancelling", "等待停止完成"), ("completed", "查看当前正式结果"), ("tampered", "初稿文件完整性校验未通过")])
def test_acceptance_errors_explain_specific_recovery(draft_task, state, message):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from src.api.auth import get_current_user
    from src.api.routes import semantic_workspace as routes
    store, root, draft = draft_task
    if state == "tampered":
        (root / "drafts" / draft["draft_id"] / "draft.json").write_text("{}", encoding="utf-8")
    else:
        store.update_semantic_workspace_task("a", "t", status=state)
    app = FastAPI()
    app.include_router(routes.router)
    app.dependency_overrides[get_current_user] = lambda: {"user_id": "a", "execution_generation": 0}
    with TestClient(app) as client:
        response = client.post("/api/semantic-workspace/tasks/t/draft/accept", json={
            "expected_revision": 1, "draft_id": draft["draft_id"], "accept_unverified": True,
        })
        assert response.status_code == 409
        assert message in response.json()["detail"]
        assert str(root) not in response.text


@pytest.mark.asyncio
async def test_publication_interruption_resumes_same_frozen_acceptance(draft_task, monkeypatch):
    from src.delivery_publishing.repository import DeliveryPublishingRepository
    store, root, draft = draft_task
    commit = DeliveryPublishingRepository.commit_delivery
    calls = 0
    def interrupt_once(self, *args):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("合成提交中断")
        return commit(self, *args)
    monkeypatch.setattr(DeliveryPublishingRepository, "commit_delivery", interrupt_once)
    params = dict(store=store, manager=None, output_root=root.parent, owner_id="a", task_id="t",
                  source_revision=1, draft_id=draft["draft_id"])
    with pytest.raises(RuntimeError, match="合成提交中断"):
        await accept_draft(**params)
    assert store.get_semantic_workspace_task("a", "t")["status"] == "candidate_ready"
    result = await accept_draft(**params)
    assert result["revision"] == 2
    assert store.get_semantic_workspace_revision("a", "t", 3) is None
    assert (await accept_draft(**params)) == result
