"""上传工作簿预览合同：样式、窗口与权限均使用合成文件。"""
import io
import os
import subprocess
from pathlib import Path
from zipfile import ZipFile
import pytest
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from tests.test_data_source_upload_api import _make_client


def test_workbook_preview_preserves_cells_and_all_sheets(tmp_path, monkeypatch):
    client = _make_client(tmp_path, monkeypatch, max_bytes=100000)
    book = Workbook()
    sheet = book.active
    sheet.title = "明细"
    sheet["A1"] = "标题"
    sheet.merge_cells("A1:C1")
    sheet["A1"].font = Font(bold=True)
    sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet["B2"] = datetime(2026, 3, 5)
    sheet["B2"].number_format = "yyyy-mm-dd"
    sheet["AD101"] = "最后一格"
    sheet["C2"] = "=SUM(1,2)"
    sheet["D2"] = "=1+1"
    sheet.merge_cells("D2:E2")
    book.create_sheet("汇总")["A1"] = "合计"
    buffer = io.BytesIO()
    book.save(buffer)
    upload = client.post("/api/data-sources/uploads", files={"file": ("sample.xlsx", buffer.getvalue())})
    assert upload.status_code == 200, upload.text
    url = f"/api/data-sources/uploads/{upload.json()['upload_id']}/workbook-preview"
    response = client.get(url)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["sheets"] == ["明细", "汇总"]
    assert data["total_rows"] == 101 and data["total_columns"] == 30
    assert data["cells"][0][0]["text"] == "标题"
    assert data["cells"][0][0]["style"]["fontWeight"] == "bold"
    assert data["cells"][1][1]["text"] == "2026-03-05"
    assert data["cells"][1][2]["text"] == "=SUM(1,2)（结果未缓存）"
    assert data["cells"][1][3]["text"] == "=1+1（结果未缓存）"
    assert {"row": 0, "column": 0, "rows": 1, "columns": 3} in data["merges"]
    last = client.get(url, params={"row": 100, "column": 29}).json()
    assert last["cells"][0][0]["text"] == "最后一格"
    assert client.get(url, params={"sheet": 1}).json()["cells"][0][0]["text"] == "合计"
    assert client.get(url, params={"row": -1}).status_code == 422
    other = _make_client(tmp_path, monkeypatch, user_id="user-b")
    assert other.get(url).status_code == 404


def test_legacy_xls_merged_title_survives_window_boundary(tmp_path, monkeypatch):
    converter = Path(os.environ.get("ProgramFiles", "")) / "LibreOffice/program/soffice.com"
    if not converter.is_file():
        pytest.skip("合成 XLS 固件需要本机 LibreOffice")
    book = Workbook()
    book.active["A1"] = "跨窗口标题"
    book.active.merge_cells("A1:A101")
    book.active["B101"] = "最后一行"
    source = tmp_path / "merged.xlsx"
    book.save(source)
    # 输入完全由测试生成；宿主工具仅用于构造旧二进制格式测试固件。
    subprocess.run([str(converter), f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}", "--headless", "--convert-to", "xls", "--outdir", str(tmp_path), str(source)], check=True, capture_output=True, timeout=45)
    client = _make_client(tmp_path / "api", monkeypatch, max_bytes=100000)
    upload = client.post("/api/data-sources/uploads", files={"file": ("merged.xls", (tmp_path / "merged.xls").read_bytes())})
    assert upload.status_code == 200, upload.text
    response = client.get(f"/api/data-sources/uploads/{upload.json()['upload_id']}/workbook-preview", params={"row": 100})
    assert response.status_code == 200, response.text
    assert response.json()["cells"][0][0]["text"] == "跨窗口标题"


def test_workbook_rejects_excessive_merge_before_rendering(tmp_path, monkeypatch):
    book = Workbook()
    book.active["A1"] = "巨大合并区"
    book.active.merge_cells("A1:A2")
    raw = io.BytesIO()
    book.save(raw)
    patched = io.BytesIO()
    with ZipFile(raw) as source, ZipFile(patched, "w") as target:
        for entry in source.infolist():
            body = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                body = body.replace(b'ref="A1:A2"', b'ref="A1:A100001"')
            target.writestr(entry, body)
    client = _make_client(tmp_path, monkeypatch, max_bytes=100000)
    upload = client.post("/api/data-sources/uploads", files={"file": ("merge.xlsx", patched.getvalue())})
    result = client.get(f"/api/data-sources/uploads/{upload.json()['upload_id']}/workbook-preview")
    assert result.status_code == 422
    assert "合并范围超过" in result.json()["detail"]
