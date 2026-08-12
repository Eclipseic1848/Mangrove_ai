# -*- coding: utf-8 -*-
"""Phase 4B 批次 2：来源检查器的确定性与证据门禁。"""
from __future__ import annotations

import hashlib
import json

from openpyxl import Workbook
import pyarrow as pa
import pyarrow.parquet as pq

from src.data_prep.document_models import DocumentElement, ElementType
from src.semantic_harness.inspection_models import InspectionStatus, TargetKind
from src.semantic_harness.inspectors.document import inspect_document_elements
from src.semantic_harness.inspectors.tabular import inspect_tabular_path


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _inspect(path, artifact_id="artifact_a"):
    data = path.read_bytes()
    return inspect_tabular_path(
        artifact_id=artifact_id,
        artifact_sha256=_sha(data),
        path=path,
        original_name=path.name,
        declared_media_type="application/octet-stream",
    )


def test_csv_inspector_profiles_columns_and_masks_sensitive_samples(tmp_path):
    path = tmp_path / "workload.csv"
    path.write_text(
        "姓名,手机号,核销工作量天数,工作量费用\n"
        "谢超群,13812345678,0.5,1200\n"
        "李四,13987654321,1,2400\n",
        encoding="utf-8",
    )

    report = _inspect(path)

    assert report.status == InspectionStatus.READY
    assert report.tables[0].sampled_rows == 2
    assert [item.raw_name for item in report.tables[0].columns] == [
        "姓名",
        "手机号",
        "核销工作量天数",
        "工作量费用",
    ]
    assert report.tables[0].columns[1].sample_values == (
        "138***5678",
        "139***4321",
    )
    assert report.canonical_hash() == report.model_copy(
        update={"generated_at": report.generated_at}
    ).canonical_hash()


def test_duplicate_headers_are_preserved_as_distinct_physical_refs(tmp_path):
    path = tmp_path / "duplicate.csv"
    path.write_text("姓名,费用,费用\n谢超群,100,200\n", encoding="utf-8")

    report = _inspect(path)
    columns = report.tables[0].columns

    assert columns[1].duplicate_group == "费用"
    assert columns[2].duplicate_group == "费用"
    assert columns[1].physical_ref != columns[2].physical_ref


def test_xlsx_and_parquet_use_mature_read_only_parsers(tmp_path):
    xlsx = tmp_path / "multi.xlsx"
    workbook = Workbook()
    workbook.active.title = "明细"
    workbook.active.append(["姓名", "费用"])
    workbook.active.append(["谢超群", 100])
    workbook.create_sheet("空表")
    workbook.save(xlsx)

    xlsx_report = _inspect(xlsx)
    assert xlsx_report.status == InspectionStatus.READY
    assert [table.name for table in xlsx_report.tables] == ["明细"]
    assert any(item.code == "empty_sheet" for item in xlsx_report.diagnostics)

    parquet = tmp_path / "records.parquet"
    pq.write_table(
        pa.table({"姓名": ["谢超群"], "费用": [100]}),
        parquet,
    )
    parquet_report = _inspect(parquet)
    assert parquet_report.status == InspectionStatus.READY
    assert parquet_report.tables[0].estimated_rows == 1


def test_json_and_corrupt_input_are_classified_without_stack_trace(tmp_path):
    good = tmp_path / "records.json"
    good.write_text(
        json.dumps([{"姓名": "谢超群", "费用": 100}], ensure_ascii=False),
        encoding="utf-8",
    )
    assert _inspect(good).status == InspectionStatus.READY

    bad = tmp_path / "bad.json"
    bad.write_text("{", encoding="utf-8")
    report = _inspect(bad)
    assert report.status == InspectionStatus.CORRUPT
    assert report.diagnostics[0].code == "parse_failed"
    assert "Traceback" not in report.diagnostics[0].message


def test_document_inspector_keeps_structural_evidence_and_marks_weak_target():
    elements = (
        DocumentElement(
            element_id="section_1",
            artifact_id="doc_a",
            page=1,
            element_type=ElementType.SECTION,
            text="商务条款",
            reading_order=1,
            extractor="python-docx",
            extractor_version="1",
            metadata={"paragraph_index": 3},
        ),
        DocumentElement(
            element_id="paragraph_2",
            artifact_id="doc_a",
            page=1,
            element_type=ElementType.PARAGRAPH,
            text="付款条件为验收后30日内支付。",
            reading_order=2,
            extractor="mineru",
            extractor_version="3.4.4",
            confidence=0.95,
        ),
    )

    report = inspect_document_elements(
        artifact_id="doc_a",
        artifact_sha256="a" * 64,
        original_name="合同.docx",
        declared_media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=100,
        elements=elements,
    )

    assert report.status == InspectionStatus.READY
    assert report.document_targets[0].target_kind == TargetKind.DOCUMENT_SECTION
    assert report.document_targets[0].evidence_ready is True
    assert report.document_targets[1].evidence_ready is False
    assert any(item.code == "missing_source_position" for item in report.diagnostics)
