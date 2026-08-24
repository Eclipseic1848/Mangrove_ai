# -*- coding: utf-8 -*-
"""Pi 候选进入独立验证门的公共行为测试。"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from src.agentic_runtime.candidate_qa import inspect_candidates
from src.agentic_runtime.candidate_verifier import (
    BrokerSemanticJudge,
    CandidateVerifier,
    _SEMANTIC_JUDGE_MAX_RETRIES,
)
from src.agentic_runtime.models import (
    PermissionProfile,
    PiRuntimeRequest,
    SemanticDecision,
    SourceInput,
    TableOutputContract,
    VerificationStatus,
)
from src.model_connections import ConnectionBroker
from src.model_connections.storage import ModelConnectionRepository
from src.model_connections.vault import FernetCredentialVault


class PassingSemanticJudge:
    """模拟外部模型边界；候选的文件与来源证据仍由真实验证器检查。"""

    async def judge(self, *, objective, candidate_previews, evidence):
        assert objective
        assert "Alice" in candidate_previews[0]
        assert "Service Fee Details" in evidence[0]
        return SemanticDecision(
            passed=True,
            contains_unrequested_content=False,
            reason="候选只包含目标表格内容",
        )


class JudgeMustNotRun:
    async def judge(self, **_kwargs):
        raise AssertionError("来源证据未通过时不得调用语义模型")


class AlwaysPassingSemanticJudge:
    async def judge(self, **_kwargs):
        return SemanticDecision(
            passed=True,
            contains_unrequested_content=False,
            reason="候选满足目标",
        )


class UnavailableSemanticJudge:
    async def judge(self, **_kwargs):
        raise ValueError(
            "1 validation error for SemanticDecision Invalid JSON: EOF "
            "https://errors.pydantic.dev/2.12/v/json_invalid"
        )


def _request(
    tmp_path: Path,
    source: Path,
    *,
    objective: str | None = None,
) -> PiRuntimeRequest:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    return PiRuntimeRequest(
        user_id="user-a",
        task_id="task-a",
        revision=1,
        objective_text=objective
        or (
            "Extract Service Fee Details and output only one CSV; "
            "do not include other content."
        ),
        requested_output_formats=("csv",),
        sources=(
            SourceInput(
                upload_id="upload-a",
                original_name="contract.pdf",
                host_path=source,
                sha256=digest,
                media_type="application/pdf",
            ),
        ),
        permission_profile=PermissionProfile.STANDARD,
        model="local-model",
        base_url="http://127.0.0.1:6012/v1",
        api_key="local-runtime",
    )


def _write_pdf(path: Path) -> None:
    document = canvas.Canvas(str(path))
    document.drawString(
        72,
        760,
        "Appendix 2 - Service Fee Details - Alice - 100",
    )
    document.save()


def _write_manifest(
    output: Path,
    *,
    quote: str,
    source_name: str = "contract.pdf",
    filename: str = "service-fees.csv",
    output_format: str = "csv",
) -> None:
    (output / "candidate-manifest.json").write_text(
        json.dumps(
            {
                "version": 1,
                "artifacts": [
                    {
                        "filename": filename,
                        "format": output_format,
                        "description": "Only the requested fee table",
                        "evidence": [
                            {
                                "source": source_name,
                                "locator": "page:1",
                                "quote": quote,
                            }
                        ],
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def test_semantic_decision_normalizes_single_missing_requirement() -> None:
    """Provider 返回单条字符串缺口时，契约保留其含义而不是整次失败。"""

    decision = SemanticDecision.model_validate(
        {
            "passed": False,
            "contains_unrequested_content": False,
            "reason": "证据不足",
            "missing_requirements": "缺少完整小计证据",
        }
    )

    assert decision.missing_requirements == ["缺少完整小计证据"]


def test_semantic_judge_allows_one_bounded_structured_output_retry() -> None:
    assert _SEMANTIC_JUDGE_MAX_RETRIES == 1


@pytest.mark.asyncio
async def test_grounded_pdf_csv_candidate_passes_independent_verifier(
    tmp_path: Path,
) -> None:
    source = tmp_path / "upload-object-without-extension"
    _write_pdf(source)
    output = tmp_path / "output"
    output.mkdir()
    (output / "service-fees.csv").write_text(
        "name,fee\nAlice,100\n",
        encoding="utf-8-sig",
    )
    _write_manifest(
        output,
        quote="Service Fee Details - Alice - 100",
    )

    candidates = inspect_candidates(output, ("csv",))
    report = await CandidateVerifier(
        semantic_judge=PassingSemanticJudge()
    ).verify(
        request=_request(tmp_path, source),
        candidates=candidates,
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED
    assert report.formal_delivery_eligible is False
    assert report.evidence_count == 1
    assert {check.code for check in report.checks if check.passed} >= {
        "artifact_set",
        "source_grounding",
        "semantic_goal",
    }


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("output_format", "filename"),
    (("csv", "service-fees.csv"),
     ("json", "service-fees.json"),
     ("xlsx", "service-fees.xlsx")),
)
async def test_table_output_contract_mismatch_fails_before_semantic_judge(
    tmp_path: Path,
    output_format: str,
    filename: str,
) -> None:
    source = tmp_path / "upload-object-without-extension"
    _write_pdf(source)
    output = tmp_path / "output"
    output.mkdir()
    candidate_path = output / filename
    if output_format == "csv":
        candidate_path.write_text(
            "fee,name\n100,Alice\n",
            encoding="utf-8-sig",
        )
    elif output_format == "json":
        candidate_path.write_text(
            json.dumps([{"fee": 100, "name": "Alice"}], ensure_ascii=False),
            encoding="utf-8",
        )
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["fee", "name"])
        sheet.append([100, "Alice"])
        workbook.save(candidate_path)
    _write_manifest(
        output,
        quote="Service Fee Details - Alice - 100",
        filename=filename,
        output_format=output_format,
    )
    contract = TableOutputContract(
        format=output_format,
        exact_columns=("name", "fee"),
        json_shape="records" if output_format == "json" else None,
    )
    request = _request(tmp_path, source).model_copy(
        update={
            "requested_output_formats": (output_format,),
            "table_output_contracts": (contract,),
        }
    )

    report = await CandidateVerifier(
        semantic_judge=JudgeMustNotRun()
    ).verify(
        request=request,
        candidates=inspect_candidates(output, (output_format,)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.FAILED
    contract_check = next(
        check for check in report.checks if check.code == "table_output_contract"
    )
    assert contract_check.passed is False


@pytest.mark.asyncio
async def test_scanned_pdf_uses_authoritative_reader_for_upload_id_manifest(
    tmp_path: Path,
) -> None:
    """扫描 PDF 没有文本层时，验证器复用任务隔离的权威读取结果。"""

    source = tmp_path / "upload-object-without-extension"
    _write_pdf(source)
    output = tmp_path / "output"
    output.mkdir()
    (output / "service-fees.csv").write_text(
        "name,fee\nAlice,100\n",
        encoding="utf-8-sig",
    )
    _write_manifest(
        output,
        quote="Authoritative OCR Service Fee Details - Alice - 100",
        source_name="upload-a",
    )
    calls: list[tuple[str, str]] = []

    async def authoritative_reader(
        source_input: SourceInput,
        locator: str,
    ) -> str:
        calls.append((source_input.upload_id, locator))
        return "Authoritative OCR Service Fee Details - Alice - 100"

    report = await CandidateVerifier(
        semantic_judge=PassingSemanticJudge(),
        authoritative_reader=authoritative_reader,
    ).verify(
        request=_request(tmp_path, source),
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED
    assert calls == [("upload-a", "page:1")]


@pytest.mark.asyncio
async def test_external_verifier_uses_separate_grant_and_records_usage(
    tmp_path: Path,
) -> None:
    provider_secret = "verifier-provider-secret-9911"
    seen: dict[str, object] = {}

    def provider(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.read())
        if body.get("max_tokens") == 16:
            return httpx.Response(
                200,
                json={
                    "choices": [
                        {"message": {"role": "assistant", "content": "OK"}}
                    ]
                },
            )
        seen["authorization"] = request.headers.get("authorization")
        seen["body"] = body
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "role": "assistant",
                            "content": json.dumps(
                                {
                                    "passed": True,
                                    "contains_unrequested_content": False,
                                    "reason": "候选只包含目标数据",
                                    "missing_requirements": [],
                                },
                                ensure_ascii=False,
                            ),
                        }
                    }
                ],
                "usage": {
                    "prompt_tokens": 11,
                    "completion_tokens": 5,
                    "total_tokens": 16,
                },
            },
        )

    broker = ConnectionBroker(
        repository=ModelConnectionRepository(str(tmp_path / "webui.db")),
        vault=FernetCredentialVault.generate(),
        transport=httpx.MockTransport(provider),
        resolver=lambda _host: ["8.8.8.8"],
    )
    connection = await broker.configure_personal(
        owner_user_id="user-a",
        preset_id="deepseek",
        api_key=provider_secret,
        model="deepseek-v4-pro",
    )
    binding = broker.freeze_connection(
        "user-a",
        str(connection["connection_id"]),
    )
    source = tmp_path / "source-object"
    source.write_text(
        "name,fee\nAlice,100\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    output.mkdir()
    (output / "service-fees.csv").write_text(
        "name,fee\nAlice,100\n",
        encoding="utf-8-sig",
    )
    _write_manifest(
        output,
        quote="Alice,100",
        source_name="contract.csv",
    )
    request = PiRuntimeRequest(
        user_id="user-a",
        task_id="task-external-verify",
        revision=1,
        objective_text="只输出一份服务费用 CSV",
        requested_output_formats=("csv",),
        sources=(
            SourceInput(
                upload_id="upload-a",
                original_name="contract.csv",
                host_path=source,
                sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                media_type="text/csv",
            ),
        ),
        model_connection_id=str(connection["connection_id"]),
        model_connection_version=binding.connection_version,
        model_connection_model=binding.model,
        external_api_confirmed=True,
    )
    report = await CandidateVerifier(
        semantic_judge=BrokerSemanticJudge(
            broker=broker,
            owner_user_id="user-a",
            connection_id=str(connection["connection_id"]),
                connection_version=binding.connection_version,
                model_id=binding.model,
                task_id=request.task_id,
            revision=1,
            run_id="pi_run_external_verify",
        )
    ).verify(
        request=request,
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED
    assert seen["authorization"] == f"Bearer {provider_secret}"
    outbound = json.dumps(seen["body"], ensure_ascii=False)
    assert "只输出一份服务费用 CSV" in outbound
    assert "Alice,100" in outbound
    assert provider_secret not in outbound
    assert broker.list_usage(
        "user-a",
        task_id=request.task_id,
        revision=1,
    ) == [
        {
            "purpose": "candidate_verify",
            "status": "recorded",
            "input_tokens": 11,
            "output_tokens": 5,
            "total_tokens": 16,
            "request_count": 1,
        }
    ]


@pytest.mark.asyncio
async def test_external_verifier_retries_one_empty_semantic_response(
    tmp_path: Path,
) -> None:
    provider_secret = "verifier-provider-secret-retry"
    verify_requests = 0

    def provider(request: httpx.Request) -> httpx.Response:
        nonlocal verify_requests
        body = json.loads(request.read())
        if body.get("max_tokens") == 16:
            return httpx.Response(
                200,
                json={
                    "choices": [
                        {"message": {"role": "assistant", "content": "OK"}}
                    ]
                },
            )
        verify_requests += 1
        content = ""
        if verify_requests == 2:
            content = json.dumps(
                {
                    "passed": True,
                    "contains_unrequested_content": False,
                    "reason": "候选只包含目标数据",
                    "missing_requirements": [],
                },
                ensure_ascii=False,
            )
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"message": {"role": "assistant", "content": content}}
                ],
                "usage": {
                    "prompt_tokens": 11,
                    "completion_tokens": 5,
                    "total_tokens": 16,
                },
            },
        )

    broker = ConnectionBroker(
        repository=ModelConnectionRepository(str(tmp_path / "webui.db")),
        vault=FernetCredentialVault.generate(),
        transport=httpx.MockTransport(provider),
        resolver=lambda _host: ["8.8.8.8"],
    )
    connection = await broker.configure_personal(
        owner_user_id="user-a",
        preset_id="deepseek",
        api_key=provider_secret,
        model="deepseek-v4-pro",
    )
    binding = broker.freeze_connection(
        "user-a",
        str(connection["connection_id"]),
    )
    source = tmp_path / "source-object"
    source.write_text("name,fee\nAlice,100\n", encoding="utf-8")
    output = tmp_path / "output"
    output.mkdir()
    (output / "service-fees.csv").write_text(
        "name,fee\nAlice,100\n",
        encoding="utf-8-sig",
    )
    _write_manifest(
        output,
        quote="Alice,100",
        source_name="contract.csv",
    )
    request = PiRuntimeRequest(
        user_id="user-a",
        task_id="task-external-retry",
        revision=1,
        objective_text="只输出一份服务费用 CSV",
        requested_output_formats=("csv",),
        sources=(
            SourceInput(
                upload_id="upload-a",
                original_name="contract.csv",
                host_path=source,
                sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                media_type="text/csv",
            ),
        ),
        model_connection_id=str(connection["connection_id"]),
        model_connection_version=binding.connection_version,
        model_connection_model=binding.model,
        external_api_confirmed=True,
    )

    report = await CandidateVerifier(
        semantic_judge=BrokerSemanticJudge(
            broker=broker,
            owner_user_id="user-a",
            connection_id=str(connection["connection_id"]),
            connection_version=binding.connection_version,
            model_id=binding.model,
            task_id=request.task_id,
            revision=1,
            run_id="pi_run_external_retry",
        )
    ).verify(
        request=request,
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED
    assert verify_requests == 2


@pytest.mark.asyncio
async def test_semantic_verifier_failure_does_not_expose_validation_stack(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source-object"
    source.write_text("name,fee\nAlice,100\n", encoding="utf-8")
    output = tmp_path / "output"
    output.mkdir()
    (output / "service-fees.csv").write_text(
        "name,fee\nAlice,100\n",
        encoding="utf-8-sig",
    )
    _write_manifest(
        output,
        quote="Alice,100",
        source_name="contract.csv",
    )
    request = PiRuntimeRequest(
        user_id="user-a",
        task_id="task-safe-error",
        revision=1,
        objective_text="只输出一份服务费用 CSV",
        requested_output_formats=("csv",),
        sources=(
            SourceInput(
                upload_id="upload-a",
                original_name="contract.csv",
                host_path=source,
                sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                media_type="text/csv",
            ),
        ),
        model="local-model",
        base_url="http://127.0.0.1:6012/v1",
        api_key="local-runtime",
    )

    report = await CandidateVerifier(
        semantic_judge=UnavailableSemanticJudge()
    ).verify(
        request=request,
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    semantic_check = next(
        check for check in report.checks if check.code == "semantic_goal"
    )
    assert report.status is VerificationStatus.INCONCLUSIVE
    assert semantic_check.summary == (
        "语义验证服务暂时不可用，请稍后重新验证候选。"
    )
    assert "pydantic" not in semantic_check.summary.lower()


@pytest.mark.asyncio
async def test_retry_semantic_verification_reuses_grounded_evidence_only(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source-object"
    source.write_text("name,fee\nAlice,100\n", encoding="utf-8")
    output = tmp_path / "output"
    output.mkdir()
    (output / "service-fees.csv").write_text(
        "name,fee\nAlice,100\n",
        encoding="utf-8-sig",
    )
    _write_manifest(
        output,
        quote="Alice,100",
        source_name="contract.csv",
    )
    request = PiRuntimeRequest(
        user_id="user-a",
        task_id="task-semantic-retry-only",
        revision=1,
        objective_text="只输出一份服务费用 CSV",
        requested_output_formats=("csv",),
        sources=(
            SourceInput(
                upload_id="upload-a",
                original_name="contract.csv",
                host_path=source,
                sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                media_type="text/csv",
            ),
        ),
        model="local-model",
        base_url="http://127.0.0.1:6012/v1",
        api_key="local-runtime",
    )
    candidates = inspect_candidates(output, ("csv",))
    previous = await CandidateVerifier(
        semantic_judge=UnavailableSemanticJudge()
    ).verify(
        request=request,
        candidates=candidates,
        manifest_path=output / "candidate-manifest.json",
    )
    source.unlink()

    report = await CandidateVerifier(
        semantic_judge=AlwaysPassingSemanticJudge()
    ).retry_semantic_verification(
        request=request,
        candidates=candidates,
        manifest_path=output / "candidate-manifest.json",
        previous_report=previous,
    )

    assert report.status is VerificationStatus.PASSED
    assert report.evidence_count == 1
    assert [check.code for check in report.checks] == [
        "artifact_set",
        "artifact_count",
        "source_grounding",
        "semantic_goal",
    ]


@pytest.mark.asyncio
async def test_unmatched_source_quote_fails_closed_before_semantic_judge(
    tmp_path: Path,
) -> None:
    source = tmp_path / "upload-object-without-extension"
    _write_pdf(source)
    output = tmp_path / "output"
    output.mkdir()
    (output / "service-fees.csv").write_text(
        "name,fee\nMallory,999\n",
        encoding="utf-8-sig",
    )
    _write_manifest(
        output,
        quote="Mallory - 999",
    )

    candidates = inspect_candidates(output, ("csv",))
    report = await CandidateVerifier(
        semantic_judge=JudgeMustNotRun()
    ).verify(
        request=_request(tmp_path, source),
        candidates=candidates,
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.FAILED
    assert report.formal_delivery_eligible is False
    failed = {check.code for check in report.checks if not check.passed}
    assert "source_grounding" in failed


@pytest.mark.asyncio
async def test_explicit_csv_and_json_goal_accepts_two_candidates(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.csv"
    source.write_text(
        "order_id,region,amount\n"
        "SYN-001,华东,120.50\n"
        "SYN-002,华南,200.00\n"
        "SYN-003,华东,80.00\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    output.mkdir()
    (output / "result.csv").write_text(
        "region,total_amount\n华东,200.50\n华南,200.00\n",
        encoding="utf-8",
    )
    (output / "result.json").write_text(
        json.dumps(
            [
                {"region": "华东", "total_amount": 200.50},
                {"region": "华南", "total_amount": 200.00},
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (output / "candidate-manifest.json").write_text(
        json.dumps(
            {
                "version": 1,
                "artifacts": [
                    {
                        "filename": filename,
                        "format": output_format,
                        "description": description,
                        "evidence": [
                            {
                                "source": "source.csv",
                                "locator": "row=2",
                                "quote": "SYN-001,华东,120.50",
                            }
                        ],
                    }
                    for filename, output_format, description in (
                        ("result.csv", "csv", "按区域汇总的 CSV"),
                        ("result.json", "json", "按区域汇总的 JSON"),
                    )
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    request = PiRuntimeRequest(
        user_id="user-a",
        task_id="task-multi-output",
        revision=1,
        objective_text=(
            "仅处理这个纯合成 CSV。按 region 汇总 amount，金额保留两位小数；"
            "输出一份 CSV 和一份 JSON。CSV 列固定为 region,total_amount，"
            "按 region 升序；JSON 使用相同字段和值。不要补造任何记录。"
        ),
        requested_output_formats=("csv", "json"),
        sources=(
            SourceInput(
                upload_id="upload-synthetic",
                original_name="source.csv",
                host_path=source,
                sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                media_type="text/csv",
            ),
        ),
        model="local-model",
        base_url="http://127.0.0.1:6012/v1",
        api_key="local-runtime",
    )

    report = await CandidateVerifier(
        semantic_judge=AlwaysPassingSemanticJudge(),
    ).verify(
        request=request,
        candidates=inspect_candidates(output, ("csv", "json")),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED
    assert all(check.passed for check in report.checks)


@pytest.mark.asyncio
async def test_explicit_single_file_goal_rejects_multiple_candidates(
    tmp_path: Path,
) -> None:
    source = tmp_path / "upload-object-without-extension"
    _write_pdf(source)
    output = tmp_path / "output"
    output.mkdir()
    for name in ("part-1.csv", "part-2.csv"):
        (output / name).write_text(
            "name,fee\nAlice,100\n",
            encoding="utf-8-sig",
        )
    (output / "candidate-manifest.json").write_text(
        json.dumps(
            {
                "version": 1,
                "artifacts": [
                    {
                        "filename": name,
                        "format": "csv",
                        "description": "fee table",
                        "evidence": [
                            {
                                "source": "contract.pdf",
                                "locator": "page:1",
                                "quote": "Service Fee Details - Alice - 100",
                            }
                        ],
                    }
                    for name in ("part-1.csv", "part-2.csv")
                ],
            }
        ),
        encoding="utf-8",
    )

    report = await CandidateVerifier(
        semantic_judge=JudgeMustNotRun()
    ).verify(
        request=_request(
            tmp_path,
            source,
            objective="抽取目标内容，输出一张单独的表，以 CSV 格式输出。",
        ),
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.FAILED
    failed = {check.code for check in report.checks if not check.passed}
    assert "artifact_count" in failed


@pytest.mark.asyncio
async def test_csv_evidence_uses_original_serialized_row(
    tmp_path: Path,
) -> None:
    source = tmp_path / "upload-object-without-extension"
    source.write_text(
        "name,fee,source_ref\nAlice,100,row=2\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    output.mkdir()
    (output / "result.csv").write_text(
        "name,fee,source_ref\nAlice,100,row=2\n",
        encoding="utf-8-sig",
    )
    (output / "candidate-manifest.json").write_text(
        json.dumps(
            {
                "version": 1,
                "artifacts": [
                    {
                        "filename": "result.csv",
                        "format": "csv",
                        "description": "requested row",
                        "evidence": [
                            {
                                "source": "contract.csv",
                                "locator": "row=2",
                                "quote": "Alice,100,row=2",
                            }
                        ],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    request = _request(tmp_path, source)
    request.sources[0].original_name = "contract.csv"
    report = await CandidateVerifier(
        semantic_judge=AlwaysPassingSemanticJudge()
    ).verify(
        request=request,
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED


@pytest.mark.asyncio
@pytest.mark.parametrize("locator", ("sheet:费用", "sheet:费用 row:2"))
async def test_xlsx_evidence_supports_extensionless_upload_object(
    tmp_path: Path,
    locator: str,
) -> None:
    source = tmp_path / "upload-object-without-extension"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "费用"
    worksheet.append(["姓名", "金额"])
    worksheet.append(["Alice", 100])
    workbook.save(source)

    output = tmp_path / "output"
    output.mkdir()
    (output / "result.csv").write_text(
        "姓名,金额\nAlice,100\n",
        encoding="utf-8-sig",
    )
    (output / "candidate-manifest.json").write_text(
        json.dumps(
            {
                "version": 1,
                "artifacts": [
                    {
                        "filename": "result.csv",
                        "format": "csv",
                        "description": "费用表中的目标行",
                        "evidence": [
                            {
                                "source": "book.xlsx",
                                "locator": locator,
                                "quote": "Alice\t100",
                            }
                        ],
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    request = _request(tmp_path, source)
    request.sources[0].original_name = "book.xlsx"
    request.sources[0].media_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    report = await CandidateVerifier(
        semantic_judge=AlwaysPassingSemanticJudge()
    ).verify(
        request=request,
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("locator", "quote"),
    (
        ("table:0 row:1", "Alice | 100"),
        (
            "table:0",
            "Table 1: 费用明细 (2 rows × 2 cols)\n"
            "姓名 | 金额\n"
            "Alice | 100",
        ),
    ),
)
async def test_docx_table_evidence_accepts_markdown_cell_separators(
    tmp_path: Path,
    locator: str,
    quote: str,
) -> None:
    source = tmp_path / "upload-object-without-extension"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "姓名"
    table.rows[0].cells[1].text = "金额"
    table.rows[1].cells[0].text = "Alice"
    table.rows[1].cells[1].text = "100"
    document.save(source)

    output = tmp_path / "output"
    output.mkdir()
    (output / "result.csv").write_text(
        "姓名,金额\nAlice,100\n",
        encoding="utf-8-sig",
    )
    (output / "candidate-manifest.json").write_text(
        json.dumps(
            {
                "version": 1,
                "artifacts": [
                    {
                        "filename": "result.csv",
                        "format": "csv",
                        "description": "目标表格行",
                        "evidence": [
                            {
                                "source": "source.docx",
                                "locator": locator,
                                "quote": quote,
                            }
                        ],
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    request = _request(tmp_path, source)
    request.sources[0].original_name = "source.docx"
    report = await CandidateVerifier(
        semantic_judge=AlwaysPassingSemanticJudge()
    ).verify(
        request=request,
        candidates=inspect_candidates(output, ("csv",)),
        manifest_path=output / "candidate-manifest.json",
    )

    assert report.status is VerificationStatus.PASSED
