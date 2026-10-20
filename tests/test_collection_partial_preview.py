"""缺少单项证据时仍可核对已证明字段，但不能把待核对资料算作完成。"""
import json
import hashlib
import pytest
from pathlib import Path

from openpyxl import load_workbook

from src.conductor.bank_benefits_export import export_snapshot
from src.conductor.collection_results import collection_outcome
from src.config.settings import settings
from tests.test_download_preview import make_client


def partial_snapshot():
    return {"schema_version": "bank-benefits-v2", "query": "采集10篇权益", "scope": {
        "target_count": 10, "strictness": "strict"}, "banks": [
        {"bank": "光大银行", "stop_reason": "time_budget"}], "candidates": [
        {"bank": "光大银行", "status": "review_required", "reason": "image_incomplete",
         "title": "私行停车权益", "content": "私行停车每月2次", "metadata": {"note_id": "n1"},
         "url": "https://www.xiaohongshu.com/explore/n1?xsec_token=secret&source=search#secret",
         "images": [{"image_index": 1, "status": "recognized", "text": "办理路径：私行专区"},
                    {"image_index": 2, "status": "failed", "text": "不能使用的失效识别内容"}],
         "assessment": {"relevant": True, "benefits": [{
             "benefit_name": "停车", "quota": "每月2次", "threshold": "资产满100万元",
             "app_path": "私行专区", "usage_period": "不能使用的失效识别内容",
             "evidence_status": "quoted_unreviewed", "missing_reasons": {}, "evidence": {
                 "benefit_name": {"source_id": "note", "quote": "停车"},
                 "quota": {"source_id": "note", "quote": "每月2次"},
                 "threshold": {"source_id": "note", "quote": "资产满100万元"},
                 "app_path": {"source_id": "image:1", "quote": "私行专区"},
                 "usage_period": {"source_id": "image:2", "quote": "不能使用的失效识别内容"},
             }}]}},
        {"bank": "光大银行", "status": "excluded", "reason": "audience_mismatch",
         "assessment": {"benefits": [{"benefit_name": "不应显示"}]}},
        {"bank": "光大银行", "status": "unprocessed", "reason": "time_budget"},
        {"bank": "光大银行", "status": "review_required", "reason": "step_timeout",
         "assessment": {"benefits": [{"benefit_name": "无证据不得显示"}]}},
    ]}


def test_partial_preview_rechecks_quotes_counts_candidates_and_preserves_original(tmp_path, monkeypatch):
    client, _, root, _ = make_client(tmp_path, monkeypatch)
    snapshot = partial_snapshot()
    path = root / "data.json"
    raw = json.dumps(snapshot, ensure_ascii=False).encode("utf-8")
    path.write_bytes(raw)
    response = client.get("/api/downloads/mine", params={"preview": "data.json"})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["candidate_counts"] == {"found": 4, "excluded": 1, "review_required": 2,
                                        "unfinished": 1, "selected": 0, "target_count": 10}
    assert data["candidate_summary"]["total_rows"] == 4
    assert data["candidate_summary"]["rows"][0]["index"] == 1
    assert data["candidate_summary"]["rows"][0]["note_url"] == "https://www.xiaohongshu.com/explore/n1?source=search"
    table = data["benefit_table"]
    assert table["total_rows"] == 1
    benefit = table["rows"][0]
    assert benefit["source_status"] == "review_required"
    assert benefit["partial_evidence"] is True and benefit["review_reason"] == "图片资料尚未识别完整"
    assert benefit["benefit_name"] == "停车" and benefit["quota"] == "每月2次" and benefit["app_path"] == "私行专区"
    assert benefit["threshold"] is None and benefit["usage_period"] is None
    assert set(benefit["evidence"]) == {"benefit_name", "quota", "app_path"}
    assert client.get("/api/downloads/mine/data.json").content == raw == path.read_bytes()


def test_candidate_summary_is_bounded_and_rejects_unsafe_links(tmp_path, monkeypatch):
    client, _, root, _ = make_client(tmp_path, monkeypatch)
    snapshot = partial_snapshot()
    snapshot["candidates"] = [{"status": "excluded", "title": "中" * 400,
                               "url": url, "reason": "unrelated"} for url in
                              ["http://127.0.0.1/private", "https://secret@example.com/note", "javascript:alert(1)"]]
    snapshot["candidates"].extend({"status": "unprocessed"} for _ in range(205))
    (root / "data.json").write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
    data = client.get("/api/downloads/mine", params={"preview": "data.json"}).json()
    assert data["candidate_counts"]["found"] == 208 and data["candidate_counts"]["unfinished"] == 205
    summary = data["candidate_summary"]
    assert summary["total_rows"] == 208 and summary["truncated"] and len(summary["rows"]) == 200
    assert all(row["note_url"] is None for row in summary["rows"][:3])
    assert len(summary["rows"][0]["source_title"]) <= 256


def test_partial_export_keeps_review_rows_separate_and_never_reaches_strict_target(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_prep_artifact_root", str(tmp_path))
    snapshot = partial_snapshot()
    image_path = tmp_path / "partial/images/1.png"
    image_path.parent.mkdir(parents=True)
    image_path.write_bytes(b"frozen-image")
    snapshot["candidates"][0]["images"][0].update(
        artifact_path=str(image_path), sha256=hashlib.sha256(image_path.read_bytes()).hexdigest())
    result = export_snapshot({"task_id": "partial", "bank_collection": snapshot})
    outcome = result["collection_outcome"]
    assert outcome["selected_count"] == 0 and not outcome["ready_for_review"] and not outcome["target_complete"]
    assert outcome["candidate_counts"]["review_required"] == 2
    raw = json.loads(Path(result["outputs"]["json"]).read_text(encoding="utf-8"))
    assert raw["candidates"] == snapshot["candidates"]
    workbook = load_workbook(result["outputs"]["xlsx"], read_only=True)
    try:
        assert len(list(workbook["权益明细"].values)) == 1
        review = list(workbook["待核对权益"].values)
        assert len(review) == 2
        row = dict(zip(review[0], review[1]))
        assert row["benefit_name"] == "停车" and row["quota"] == "每月2次"
        assert row["threshold"] is None and row["usage_period"] is None
        assert row["partial_evidence"] is True and row["review_reason"] == "图片资料尚未识别完整"
        assert row["note_url"] == "https://www.xiaohongshu.com/explore/n1?source=search"
        assert "候选明细审计" in workbook.sheetnames
    finally:
        workbook.close()


def test_partial_count_does_not_promote_strict_delivery():
    snapshot = partial_snapshot()
    snapshot["candidates"].append({"bank": "光大银行", "status": "selected"})
    outcome = collection_outcome(snapshot)
    assert outcome["selected_count"] == 1 and outcome["status"] == "incomplete"
    assert not outcome["ready_for_review"] and not outcome["target_complete"]
    assert outcome["candidate_counts"]["selected"] == 1


def test_live_topic_preview_exposes_saved_revision_without_promoting_delivery(tmp_path, monkeypatch):
    client, _, root, _ = make_client(tmp_path, monkeypatch)
    snapshot = partial_snapshot()
    snapshot.update(schema_version="evidence-collection-v1", in_progress=True, snapshot_revision=3)
    (root / "data-in-progress.json").write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
    data = client.get("/api/downloads/mine", params={"preview": "data-in-progress.json"}).json()
    assert data["in_progress"] is True and data["snapshot_revision"] == 3
    assert data["candidate_counts"]["selected"] == 0 and data["candidate_summary"]["total_rows"] == 4
    assert data["benefit_table"] is None


@pytest.mark.parametrize("code,label", [
    ("image_incomplete", "图片资料尚未识别完整"),
    ("images_incomplete", "图片资料尚未识别完整"),
    ("comments_incomplete", "评论资料尚未读取完整"),
    ("screen_outcome_unknown", "本篇初筛结果尚未确认"),
    ("extract_outcome_unknown", "本篇权益抽取结果尚未确认"),
    ("images_outcome_unknown", "本篇图片识别结果尚未确认"),
    ("comments_outcome_unknown", "本篇评论读取结果尚未确认"),
    ("discovery_outcome_unknown", "候选检索结果尚未确认"),
    ("audience_mismatch", "与本次要求的客群不符"),
])
def test_candidate_reason_uses_fixed_product_text(tmp_path, monkeypatch, code, label):
    client, _, root, _ = make_client(tmp_path, monkeypatch)
    snapshot = partial_snapshot()
    snapshot["candidates"] = [{"status": "review_required", "reason": code}]
    (root / "data.json").write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
    row = client.get("/api/downloads/mine", params={"preview": "data.json"}).json()["candidate_summary"]["rows"][0]
    assert row["reason_code"] == code and row["reason"] == label
    assert row["explanation"] is None and row["evidence"] == []


def test_exclusion_explanation_requires_real_quotes_and_bounds_all_values(tmp_path, monkeypatch):
    client, _, root, _ = make_client(tmp_path, monkeypatch)
    text = "普通客户权益" + "中" * 400
    note = {"status": "excluded", "reason": "audience_mismatch", "content": text,
            "metadata": {"publish_time": "2026-10-01T00:00:00+08:00", "comments": [
                {"comment_id": "c1", "content": "大众客户可办理"}]},
            "images": [{"image_index": 1, "status": "recognized", "text": "普通客户申请"},
                       {"image_index": 2, "status": "failed", "text": "伪识别不能引用"}],
            "assessment": {"tag_audience_mismatch": True, "rank_reason": "原文限定为普通客户" * 40,
                "judgment_evidence": {
                    "tag_audience_mismatch": {"source_id": "note", "quote": text},
                    "a": {"source_id": "image:1", "quote": "普通客户申请"},
                    "b": {"source_id": "comment:c1", "quote": "大众客户可办理"},
                    "c": {"source_id": "metadata", "quote": "2026-10-01T00:00:00+08:00"},
                    "d": {"source_id": "note", "quote": "普通客户权益"},
                    "e": {"source_id": "image:2", "quote": "伪识别不能引用"},
                    "f": {"source_id": "note", "quote": "原文不存在的条件"},
                }}}
    snapshot = partial_snapshot()
    snapshot["candidates"] = [note]
    (root / "data.json").write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
    row = client.get("/api/downloads/mine", params={"preview": "data.json"}).json()["candidate_summary"]["rows"][0]
    assert row["reason"] == "与本次要求的客群不符"
    assert len(row["explanation"]) == 120
    assert len(row["evidence"]) == 4
    assert all(len(ref["quote"]) <= 240 for ref in row["evidence"])
    assert not any(ref["source_id"] == "image:2" for ref in row["evidence"])
    note["assessment"]["judgment_evidence"] = {"tag_audience_mismatch": {"source_id": "note", "quote": "伪造条件"}}
    (root / "data.json").write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
    row = client.get("/api/downloads/mine", params={"preview": "data.json"}).json()["candidate_summary"]["rows"][0]
    assert row["explanation"] is None and row["evidence"] == []


def test_unknown_error_and_provider_logs_never_enter_candidate_projection(tmp_path, monkeypatch):
    client, _, root, _ = make_client(tmp_path, monkeypatch)
    raw_error = "provider HTTP 401 Authorization: Bearer sk-secret-test"
    snapshot = partial_snapshot()
    snapshot["candidates"] = [
        {"status": "review_required", "reason": raw_error, "assessment": {"rank_reason": raw_error}},
        {"status": "excluded", "reason": "unrelated", "content": "原文存在",
         "assessment": {"relevant": False, "rank_reason": raw_error,
                        "judgment_evidence": {"relevant": {"source_id": "note", "quote": "原文存在"}}}},
    ]
    raw = json.dumps(snapshot, ensure_ascii=False).encode("utf-8")
    (root / "data.json").write_bytes(raw)
    rows = client.get("/api/downloads/mine", params={"preview": "data.json"}).json()["candidate_summary"]["rows"]
    assert rows[0]["reason_code"] == "unknown" and rows[0]["reason"] == "尚待核对"
    assert all(row["explanation"] is None for row in rows)
    assert "sk-secret" not in json.dumps(rows)
    assert client.get("/api/downloads/mine/data.json").content == raw


def test_publication_exclusion_shows_actual_date_and_frozen_range(tmp_path, monkeypatch):
    client, _, root, _ = make_client(tmp_path, monkeypatch)
    snapshot = partial_snapshot()
    snapshot["scope"].update(publication_from="2026-01-01", publication_to="2026-10-12",
                             frozen_at="2026-10-12T09:37:22+08:00")
    snapshot["candidates"] = [{"status": "excluded", "reason": "publication_outside_scope",
                                "metadata": {"publish_time": "2025-10-12T12:00:00+08:00"}}]
    (root / "data.json").write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
    row = client.get("/api/downloads/mine", params={"preview": "data.json"}).json()["candidate_summary"]["rows"][0]
    assert "2025-10-12" in row["explanation"]
    assert "2026-01-01" in row["explanation"] and "2026-10-12" in row["explanation"]
    assert row["evidence"] == [{"source_id": "metadata", "quote": "2025-10-12T12:00:00+08:00"}]
