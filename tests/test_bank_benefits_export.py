"""初稿文件同源、可重开、长文本保真及公式安全。"""
import json

from openpyxl import load_workbook

from src.config.settings import settings
from src.conductor.bank_benefits_export import export_snapshot


def test_exports_preserve_source_and_never_generate_report(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_prep_artifact_root", str(tmp_path))
    text = "=不能执行的原文" + "中文" * 20000
    snapshot = {"schema_version": "bank-benefits-v1", "query": "权益", "scope": {}, "banks": [],
                "candidates": [{"bank": "中信", "status": "selected", "title": "=1+1", "content": text,
                                "url": "https://example.com", "metadata": {"note_id": "n"},
                                "assessment": {"benefits": []}, "images": []}]}
    result = export_snapshot({"task_id": "test-export", "bank_collection": snapshot})
    assert "report_md" not in result["outputs"]
    exported = json.loads((tmp_path / "test-export/data.json").read_text(encoding="utf-8"))
    assert exported["candidates"][0]["content"] == text
    workbook = load_workbook(tmp_path / "test-export/data.xlsx", read_only=True)
    try:
        rows = list(workbook["笔记清单"].values)
        assert "=1+1" in rows[1]
        assert all(cell.data_type != "f" for row in workbook["笔记清单"] for cell in row)
        chunks = list(workbook["长文本"].values)[1:]
        assert "".join(row[2] for row in chunks) == text
        assert {"笔记清单", "权益明细", "评论", "图片识别", "执行说明"}.issubset(workbook.sheetnames)
    finally:
        workbook.close()


def test_recovery_snapshot_keeps_previous_download_bytes(tmp_path, monkeypatch):
    from pathlib import Path
    import pytest
    monkeypatch.setattr(settings, "data_prep_artifact_root", str(tmp_path))
    state = {"task_id": "recovery-export", "bank_collection": {
        "schema_version": "bank-benefits-v1", "query": "原查询", "scope": {}, "banks": [], "candidates": []}}
    first = export_snapshot(state)
    original = {key: (Path(path), Path(path).read_bytes()) for key, path in first["outputs"].items()}
    second = export_snapshot({**state, "collection_output_attempt": "a" * 32})
    assert all(path.read_bytes() == raw for path, raw in original.values())
    assert all(Path(path).name != original[key][0].name for key, path in second["outputs"].items())
    with pytest.raises(ValueError):
        export_snapshot({**state, "collection_output_attempt": "../escape"})
