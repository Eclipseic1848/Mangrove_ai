"""同一流程覆盖非银行主题、按需读取、冻结恢复和真实导出格式。"""
import asyncio
import json

import pytest
from openpyxl import load_workbook
from src.collectors.base import CollectResult, CollectedItem
from src.conductor import evidence_collection as flow
from src.conductor.task_spec import TaskSpec


def spec_for(topic, **scope):
    return TaskSpec(intent=f"采集小红书{topic}", platforms=["小红书"], evidence_collection={
        "source_evidence": "小红书", "queries": [{"name": topic, "keyword": topic}],
        "fields": ["name", "condition"], "target_count": 1, **scope})


@pytest.mark.parametrize("topic", ["酒店会员", "商品优惠"])
def test_topics_share_discovery_screen_extraction_and_export(topic, tmp_path, monkeypatch):
    from src.conductor.nodes.collect import collect_node
    from src.conductor.nodes.output import output_node
    calls = []
    spec = spec_for(topic)
    async def collect(self, request):
        assert not request.include_comments
        calls.append("discover")
        return CollectResult(True, "synthetic", items=[CollectedItem(content=text, metadata={
            "note_id": str(i), "note_type": "normal", "image_urls": []}) for i, text in enumerate(["无关", "会员领券"])],
            coverage={"pages": [{"has_more": False, "item_count": 2}]})
    async def model(messages, **kwargs):
        payload = json.loads(messages[-1]["content"])
        relevant = "会员领券" in payload["sources"]["note"]
        calls.append(("screen" if kwargs["max_tokens"] == 8192 else "extract", relevant))
        return json.dumps({"relevant": relevant, "records": [{"name": "领券", "condition": "会员", "evidence": {
            "name": {"source_id": "note", "quote": "会员领券"}, "condition": {"source_id": "note", "quote": "会员领券"}}}]})
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    monkeypatch.setattr(flow, "achat", model)
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path))
    state = {"task_id": "topic", "task_spec": spec}
    state.update(asyncio.run(collect_node(state)))
    assert calls == ["discover", ("screen", False), ("screen", True), ("extract", True)]
    assert state["collection_outcome"]["ready_for_review"]
    result = asyncio.run(output_node(state))
    workbook = load_workbook(result["outputs"]["xlsx"], read_only=True)
    try:
        assert workbook["结构化明细"].max_row == 2
    finally:
        workbook.close()
    assert len(state["evidence_collection"]["candidates"]) == 2


def test_resume_never_repeats_completed_steps_or_unknown_inflight(tmp_path, monkeypatch):
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path / "artifacts"))
    monkeypatch.setattr(flow.settings, "semantic_execution_root", str(tmp_path / "execution"))
    monkeypatch.setattr(flow, "execution_owner", lambda: "owner-a")
    calls = []
    async def collect(self, request):
        calls.append("discover")
        return CollectResult(True, "synthetic", items=[CollectedItem(content="规则", metadata={
            "note_id": str(i), "note_type": "normal"}) for i in range(3)],
            coverage={"pages": [{"has_more": False}]})
    interrupted = False
    async def assess(note, scope, state):
        nonlocal interrupted
        key = (note["metadata"]["note_id"], state["_assessment_stage"])
        calls.append(key)
        if key == ("1", "extract") and not interrupted:
            interrupted = True
            raise asyncio.CancelledError()
        return {"relevant": True, "records": []}
    async def images(*args):
        raise AssertionError("没有图片不能启动OCR")
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    spec = spec_for("服务规则", target_count=3)
    state = {"task_id": "resume", "task_spec": spec}
    def run():
        return asyncio.run(flow.collect_evidence(state, scope=spec.evidence_collection,
            queries=spec.evidence_collection.queries, assess=assess, exclude=flow.exclude_topic, read_images=images))
    with pytest.raises(asyncio.CancelledError):
        run()
    result = run()
    assert calls.count("discover") == 1
    assert calls.count(("0", "extract")) == 1
    assert calls.count(("1", "extract")) == 1
    assert calls.count(("2", "extract")) == 1
    assert [n["status"] for n in result["evidence_collection"]["candidates"]] == ["selected", "review_required", "selected"]
    state["model"] = "changed-model"
    with pytest.raises(ValueError, match="版本不一致"):
        run()


def test_export_separates_selected_records_from_audit(tmp_path, monkeypatch):
    from src.conductor.bank_benefits_export import export_snapshot
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path))
    snapshot = {"query": "权益", "scope": {"target_count": 1},
        "banks": [{"bank": "示例", "selected_count": 1, "stop_reason": "target_reached"}],
        "candidates": [{"bank": "示例", "status": status, "metadata": {"note_id": str(i),
                            "comments": [{"comment_id": "c", "content": "评论"}]},
                        "images": [{"image_index": 1, "status": "failed"}],
                        "assessment": {"benefits": [{"benefit_name": status}]}}
                       for i, status in enumerate(["selected", "excluded", "unprocessed"])]}
    result = export_snapshot({"task_id": "projection", "bank_collection": snapshot})
    workbook = load_workbook(result["outputs"]["xlsx"], read_only=True)
    try:
        assert workbook["权益明细"].max_row == 2
        assert workbook["候选明细审计"].max_row == 4
        assert "source_status" in next(workbook["权益明细"].values)
        for name in ("评论", "图片识别"):
            rows = list(workbook[name].values)
            column = rows[0].index("source_status")
            assert [row[column] for row in rows[1:]] == ["selected", "excluded", "unprocessed"]
    finally:
        workbook.close()


def test_private_source_access_is_owner_task_bound_and_never_public(monkeypatch):
    from src.collectors import source_access
    from src.model_connections.vault import FernetCredentialVault
    vault = FernetCredentialVault.generate()
    monkeypatch.setattr(source_access, "load_vault", lambda _: vault)
    monkeypatch.setattr(source_access, "execution_owner", lambda: "owner-a")
    url = "https://www.xiaohongshu.com/explore/note?xsec_token=synthetic-secret"
    handle = source_access.seal_access("task", "note", url)
    assert "synthetic-secret" not in handle
    item = CollectedItem(metadata={"note_id": "note"}, access_handle=handle)
    assert "access_handle" not in item.to_dict()
    assert source_access.open_access(handle, "task", "note") == url
    with pytest.raises(PermissionError):
        source_access.open_access(handle, "other", "note")
    monkeypatch.setattr(source_access, "execution_owner", lambda: "owner-b")
    with pytest.raises(PermissionError):
        source_access.open_access(handle, "task", "note")


@pytest.mark.parametrize("interrupt_discovery", [False, True])
def test_discovery_budget_counts_duplicates_and_never_replays_unknown(tmp_path, monkeypatch, interrupt_discovery):
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path / "artifacts"))
    monkeypatch.setattr(flow.settings, "semantic_execution_root", str(tmp_path / "execution"))
    monkeypatch.setattr(flow, "execution_owner", lambda: "owner")
    calls = []
    async def collect(self, request):
        calls.append(request.max_items)
        if interrupt_discovery:
            raise asyncio.CancelledError()
        return CollectResult(True, "synthetic", items=[CollectedItem(metadata={
            "note_id": str(i), "note_type": "normal"}) for i in [0, len(calls)]],
            coverage={"pages": [{"has_more": True}]})
    async def assess(*args):
        return {"relevant": False}
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    spec = spec_for("规则", initial_candidates=2, candidate_limit=4)
    def run():
        return asyncio.run(flow.collect_evidence({"task_id": "budget", "task_spec": spec},
            scope=spec.evidence_collection, queries=spec.evidence_collection.queries,
            assess=assess, exclude=flow.exclude_topic, read_images=None))
    if interrupt_discovery:
        with pytest.raises(asyncio.CancelledError):
            run()
    result = run()["evidence_collection"]["groups"][0]
    assert calls == ([2] if interrupt_discovery else [2, 2])
    assert result["candidate_attempted"] == sum(calls)
    assert result["stop_reason"] == ("interrupted_discovery_outcome_unknown" if interrupt_discovery else "candidate_limit")


@pytest.mark.parametrize("tamper", ["image", "note"])
def test_resume_rejects_changed_frozen_source_and_image(tmp_path, monkeypatch, tamper):
    from src.conductor.bank_benefits_export import export_snapshot
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path / "artifacts"))
    monkeypatch.setattr(flow.settings, "semantic_execution_root", str(tmp_path / "execution"))
    monkeypatch.setattr(flow, "execution_owner", lambda: "owner")
    async def collect(self, request):
        return CollectResult(True, "synthetic", items=[CollectedItem(content="规则", metadata={
            "note_id": "n", "note_type": "normal", "image_urls": ["https://example.com/image"]})])
    async def assess(*args):
        return {"relevant": True}
    async def images(note, task_id, raw_only):
        artifact = flow.ArtifactStore().write_raw(task_id, "n", b"synthetic", uri="image", media_type="image/png")
        return [{"image_index": 1, "status": "recognized", "sha256": artifact.sha256,
                 "artifact_path": artifact.storage_path, "text": "规则"}]
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    spec = spec_for("规则")
    state = {"task_id": "tamper", "task_spec": spec}
    def run():
        return asyncio.run(flow.collect_evidence(state, scope=spec.evidence_collection,
            queries=spec.evidence_collection.queries, assess=assess, exclude=flow.exclude_topic, read_images=images))
    result = run()
    note = result["evidence_collection"]["candidates"][0]
    path = note["images"][0]["artifact_path"] if tamper == "image" else note["source_artifact"]["storage_path"]
    flow.ArtifactStore().resolve_path(path).write_bytes(b"changed")
    with pytest.raises(ValueError, match="变化"):
        run()
    if tamper == "image":
        with pytest.raises(ValueError, match="变化"):
            export_snapshot({**state, **result})


def test_comments_are_read_only_after_scope_screening(monkeypatch):
    from src.collectors import source_access
    calls = []
    async def collect(self, request):
        calls.append((request.include_comments, request.urls))
        if request.urls:
            return CollectResult(True, "synthetic", items=[CollectedItem(metadata={"note_id": "1",
                "comments": [{"comment_id": "c", "content": "领取条件"}], "comment_coverage": {"truncated": False}})])
        return CollectResult(True, "synthetic", items=[CollectedItem(metadata={"note_id": str(i), "note_type": "normal"}) for i in range(2)])
    async def assess(note, scope, state):
        if state["_assessment_stage"] == "extract":
            assert flow.source_context(note)["comment:c"] == "领取条件"
        return {"relevant": note["metadata"]["note_id"] == "1"}
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    monkeypatch.setattr(source_access, "open_access", lambda handle, task_id, note_id: "https://www.xiaohongshu.com/explore/" + note_id)
    spec = spec_for("服务规则").model_copy(update={"include_comments": True})
    result = asyncio.run(flow.collect_evidence({"task_id": "comments", "task_spec": spec},
        scope=spec.evidence_collection, queries=spec.evidence_collection.queries,
        assess=assess, exclude=flow.exclude_topic, read_images=None))
    assert calls == [(False, []), (True, ["https://www.xiaohongshu.com/explore/1"])]
    assert result["collection_outcome"]["status"] == "complete"


def test_time_conditions_remain_distinct_and_publication_is_in_model_context(monkeypatch):
    from src.conductor import bank_benefits
    text = "活动期1月至3月，报名期1月，券可使用至6月"
    async def model(messages, **kwargs):
        sources = json.loads(messages[-1]["content"])["sources"]
        assert "2025-12-01" in sources["metadata"]
        fields = {"activity_period": "1月至3月", "registration_period": "1月", "usage_period": "至6月",
                  "valid_from": "1月", "valid_to": "6月"}
        return json.dumps({"relevant": True, "benefits": [{**fields,
            "evidence": {key: {"source_id": "note", "quote": text} for key in fields}}]})
    monkeypatch.setattr(bank_benefits, "achat", model)
    spec = TaskSpec(intent="历史权益", bank_benefits={"source_evidence": "小红书",
        "publication_from": "2025-01-01", "publication_to": "2025-12-31",
        "banks": [{"name": "示例", "keyword": "历史权益"}]})
    result = asyncio.run(bank_benefits.assess_note({"content": text, "metadata": {"publish_time": "2025-12-01T00:00:00+08:00"}},
        spec.bank_benefits, {"task_spec": spec}))
    row = result["benefits"][0]
    assert row["valid_from"] is None and row["valid_to"] is None
    assert row["activity_period"] == "1月至3月" and row["usage_period"] == "至6月"
    assert row["review_required"] and row["evidence_status"] == "quoted_unreviewed"


def test_public_conductor_resume_reuses_frozen_plan_and_rejects_changed_identity(monkeypatch):
    from langgraph.graph import StateGraph, START, END
    from langgraph.checkpoint.memory import InMemorySaver
    from src.conductor import graph as conductor
    from src.conductor.state import ConductorState
    from src.memory import _library_scope
    calls = []
    async def plan(state):
        calls.append("plan")
        return {"task_spec": spec_for("服务规则")}
    async def collect(state):
        calls.append("collect")
        if calls.count("collect") == 1:
            raise asyncio.CancelledError()
        return {"reply": "恢复完成"}
    builder = StateGraph(ConductorState)
    builder.add_node("plan", plan)
    builder.add_node("collect", collect)
    builder.add_edge(START, "plan")
    builder.add_edge("plan", "collect")
    builder.add_edge("collect", END)
    graph = builder.compile(checkpointer=InMemorySaver())
    async def get_graph():
        return graph
    monkeypatch.setattr(conductor, "_get_checkpoint_graph", get_graph)
    monkeypatch.setattr(flow.settings, "checkpoint_enabled", True)
    monkeypatch.setattr(_library_scope, "execution_owner", lambda: "owner-a")
    async def scenario():
        with pytest.raises(asyncio.CancelledError):
            await conductor.run_conductor("采集小红书服务规则", task_id="resume-graph")
        result = [payload async for kind, payload in conductor.astream_conductor("采集小红书服务规则", task_id="resume-graph") if kind == "final"]
        assert result[0]["reply"] == "恢复完成" and calls == ["plan", "collect", "collect"]
        monkeypatch.setattr(_library_scope, "execution_owner", lambda: "owner-b")
        with pytest.raises(ValueError, match="拒绝恢复"):
            await conductor.run_conductor("采集小红书服务规则", task_id="resume-graph")
        assert calls == ["plan", "collect", "collect"]
    asyncio.run(scenario())


@pytest.mark.parametrize("stage", ["screen", "extract"])
def test_truncation_preserves_candidate_without_retries(stage, tmp_path, monkeypatch):
    from src.conductor.nodes.collect import collect_node
    from src.model_connections.text_protocol import ModelOutputTruncatedError
    calls = []
    async def collect(self, request):
        return CollectResult(True, "synthetic", items=[CollectedItem(content="会员权益", metadata={
            "note_id": "one", "note_type": "normal", "image_urls": []})],
            coverage={"pages": [{"has_more": False}]})
    async def model(messages, **kwargs):
        actual = "screen" if kwargs["max_tokens"] == 8192 else "extract"
        calls.append(actual)
        if actual == stage:
            raise ModelOutputTruncatedError("模型输出截断")
        return '{"relevant": true, "records": []}'
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    monkeypatch.setattr(flow, "achat", model)
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path))
    result = asyncio.run(collect_node({"task_id": "truncated", "task_spec": spec_for("会员权益")}))
    note = result["evidence_collection"]["candidates"][0]
    assert note["status"] == "review_required" and note["reason"] == "model_output_truncated"
    assert result["collection_outcome"]["status"] == "incomplete"
    assert calls == (["screen"] if stage == "screen" else ["screen", "extract"])


@pytest.mark.parametrize("failure,model,expected", [
    ("truncated", "deepseek-flash", [8192, 65536, 131072]),
    ("invalid", "deepseek-flash", [8192, 65536]),
    ("always_truncated", "deepseek-flash", [8192, 65536, 131072]),
    ("timeout", "deepseek-flash", [8192, 65536]),
    ("truncated", "unknown-model", [8192, 65536]),
    ("truncated", "small-model", [8192, 65536]),
])
def test_extract_retries_only_known_truncation_with_larger_supported_budget(failure, model, expected, tmp_path, monkeypatch):
    from src.conductor.nodes.collect import collect_node
    from src.model_connections.text_protocol import ModelOutputTruncatedError
    calls = []
    if model == "small-model":
        monkeypatch.setattr(flow, "model_max_output_tokens", lambda _: 32768)
    async def collect(self, request):
        return CollectResult(True, "synthetic", items=[CollectedItem(content="会员权益", metadata={
            "note_id": "one", "note_type": "normal", "image_urls": []})],
            coverage={"pages": [{"has_more": False}]})
    async def generate(messages, **kwargs):
        calls.append(kwargs["max_tokens"])
        if len(calls) == 2 or (len(calls) > 2 and failure == "always_truncated"):
            if failure == "timeout":
                raise asyncio.TimeoutError()
            if failure in {"truncated", "always_truncated"}:
                raise ModelOutputTruncatedError("截断")
            raise ValueError("结构错误")
        return '{"relevant": true, "records": []}'
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    monkeypatch.setattr(flow, "achat", generate)
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path))
    result = asyncio.run(collect_node({"task_id": "budget-retry", "task_spec": spec_for("会员权益"), "model": model}))
    assert calls == expected
    note = result["evidence_collection"]["candidates"][0]
    assert note["status"] == ("selected" if failure == "truncated" and len(expected) == 3 else "unprocessed" if failure == "timeout" else "review_required")


@pytest.mark.parametrize("raw_only", [False, True])
def test_explicit_text_only_scope_never_reads_images(tmp_path, monkeypatch, raw_only):
    monkeypatch.setattr(flow.settings, "data_prep_artifact_root", str(tmp_path))
    spec = spec_for("设备说明", raw_only=raw_only, include_images=False, initial_candidates=1, candidate_limit=1)
    async def collect(self, request):
        return CollectResult(True, "synthetic", items=[CollectedItem(content="设备说明", metadata={
            "note_id": "one", "note_type": "normal", "image_urls": ["https://example.com/photo.jpg"]})])
    async def assess(*args):
        return {"relevant": True, "records": []}
    async def forbidden(*args):
        raise AssertionError("明确不读图，不得下载或OCR")
    monkeypatch.setattr(flow.SocialMediaCollector, "collect", collect)
    result = asyncio.run(flow.collect_evidence({"task_id": "text-only", "task_spec": spec},
        scope=spec.evidence_collection, queries=spec.evidence_collection.queries,
        assess=assess, exclude=flow.exclude_topic, read_images=forbidden))
    assert result["collection_outcome"]["selected_count"] == 1
    assert result["evidence_collection"]["candidates"][0]["images"] == []
