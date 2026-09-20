"""真实隔离容器转换合成 Office 文件，不调用模型。"""
import io
import os
from pathlib import Path
import shutil
import subprocess

import pytest
from docx import Document
from tests.test_data_source_upload_api import _make_client


def test_office_preview_rejects_non_office_file(tmp_path, monkeypatch):
    client = _make_client(tmp_path, monkeypatch)
    upload = client.post("/api/data-sources/uploads", files={"file": ("sample.txt", b"text")})
    response = client.get(f"/api/data-sources/uploads/{upload.json()['upload_id']}/office-preview")
    assert response.status_code == 415


@pytest.mark.parametrize("extension", ["docx", "pptx", "doc", "ppt"])
def test_office_preview_returns_pdf_without_mutating_original(tmp_path, monkeypatch, extension):
    if not shutil.which("docker") or subprocess.run(["docker", "image", "inspect", "mangrove/office-preview:local"], capture_output=True).returncode:
        pytest.skip("需要先构建隔离预览镜像")
    client = _make_client(tmp_path, monkeypatch, max_bytes=1000000)
    if extension.startswith("doc"):
        document = Document()
        document.add_heading("合成文档", 1)
        document.add_paragraph("预览不修改原件")
    else:
        from pptx import Presentation
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = "合成幻灯片"
        document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = "第二页"
    stream = io.BytesIO()
    document.save(stream)
    original = stream.getvalue()
    if extension in {"doc", "ppt"}:
        converter = Path(os.environ.get("ProgramFiles", "")) / "LibreOffice/program/soffice.com"
        if not converter.is_file():
            pytest.skip("生成旧格式合成固件需要本机 LibreOffice")
        source = tmp_path / f"sample.{extension}x"
        source.write_bytes(original)
        subprocess.run([str(converter), f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}", "--headless", "--convert-to", extension, "--outdir", str(tmp_path), str(source)], check=True, capture_output=True, timeout=45)
        original = (tmp_path / f"sample.{extension}").read_bytes()
    upload = client.post("/api/data-sources/uploads", files={"file": (f"sample.{extension}", original)})
    assert upload.status_code == 200, upload.text
    url = f"/api/data-sources/uploads/{upload.json()['upload_id']}"
    response = client.get(url + "/office-preview")
    assert response.status_code == 200, response.text[:200]
    assert response.content.startswith(b"%PDF-")
    assert response.headers["content-type"] == "application/pdf"
    from pypdf import PdfReader
    pages = PdfReader(io.BytesIO(response.content)).pages
    assert len(pages) == (1 if extension.startswith("doc") else 2)
    assert "合成" in pages[0].extract_text()
    assert client.get(url + "/content").content == original
