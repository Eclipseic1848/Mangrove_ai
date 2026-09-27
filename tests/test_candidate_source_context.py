"""独立语义核验须看到引用之外的冻结上下文；不伪称截断材料完整。"""
import hashlib
import json

import pytest

from src.agentic_runtime.candidate_qa import inspect_candidates
from src.agentic_runtime.candidate_verifier import CandidateVerifier
from src.agentic_runtime.models import PiRuntimeRequest, SemanticDecision, SourceInput, VerificationStatus


@pytest.mark.asyncio
@pytest.mark.parametrize("mode", ["complete", "retry", "oversize", "changed", "total_limit", "unreadable", "docx_header", "xlsx_comment"])
async def test_semantic_review_distinguishes_whole_frozen_text_from_quotes(tmp_path, mode):
    source = tmp_path / "frozen-upload"
    text = "已引用的记录\n未被引用的字段：网站丙，评分尺度尚未定义。"
    source.write_text(text + ("甲" * 4100 if mode == "oversize" else ""), encoding="utf-8")
    original_name = "source.txt"
    if mode == "docx_header":
        from docx import Document
        document = Document()
        document.add_paragraph("已引用的记录")
        document.sections[0].header.paragraphs[0].text = "网站丙"
        document.save(source)
        original_name = "source.docx"
    elif mode == "xlsx_comment":
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        workbook = Workbook()
        workbook.active["A1"] = "已引用的记录"
        workbook.active["A1"].comment = Comment("网站丙", "原文作者")
        workbook.save(source)
        workbook.close()
        original_name = "source.xlsx"
    record = SourceInput(upload_id="source", original_name=original_name, host_path=source,
                         sha256=hashlib.sha256(source.read_bytes()).hexdigest())
    if mode == "changed":
        source.write_text(text + "\n冻结后改动", encoding="utf-8")
    sources = [record]
    for index in range(5 if mode == "total_limit" else 1 if mode == "unreadable" else 0):
        extra = tmp_path / f"other-{index}.txt"
        extra.write_text("额外原文" * 900, encoding="utf-8")
        sources.append(SourceInput(upload_id=f"extra-{index}", original_name=extra.name, host_path=extra,
                                  sha256=hashlib.sha256(extra.read_bytes()).hexdigest()))
        if mode == "unreadable":
            extra.unlink()
    output = tmp_path / "output"
    output.mkdir()
    (output / "report.json").write_text('{"conclusion":"原文没有其他网站"}', encoding="utf-8")
    manifest = output / "candidate-manifest.json"
    manifest.write_text(json.dumps({"version": 1, "artifacts": [{"filename": "report.json", "format": "json",
        "description": "核对报告", "evidence": [{"source": original_name, "locator": "all", "quote": "已引用的记录"}]}]},
        ensure_ascii=False), encoding="utf-8")
    request = PiRuntimeRequest(user_id="owner", task_id="task", revision=1,
        objective_text="基于提供的原文生成报告，不编造缺失信息。", requested_output_formats=("json",),
        sources=tuple(sources), model="local", base_url="http://127.0.0.1:6012/v1", api_key="test")
    seen = []

    class Judge:
        async def judge(self, *, objective, candidate_previews, evidence):
            seen.append(evidence)
            if mode == "retry" and len(seen) == 1:
                raise ValueError("合成服务暂不可用")
            return SemanticDecision(passed=False, contains_unrequested_content=False, reason="需核对原文未引用部分")

    verifier = CandidateVerifier(semantic_judge=Judge())
    candidates = inspect_candidates(output, ("json",))
    report = await verifier.verify(request=request, candidates=candidates, manifest_path=manifest)
    if mode == "retry":
        assert report.status is VerificationStatus.INCONCLUSIVE
        report = await verifier.retry_semantic_verification(request=request, candidates=candidates,
            manifest_path=manifest, previous_report=report)
    assert report.status is VerificationStatus.FAILED
    complete = mode in {"complete", "retry"}
    assert any("FULL_FROZEN_SOURCE=" in item for item in seen[-1]) is complete
    if complete:
        assert any("网站丙" in item for item in seen[-1])
    else:
        assert all("网站丙" not in item for item in seen[-1])
        assert any("PARTIAL_SOURCE_CONTEXT" in item for item in seen[-1])
