"""有界表格内容保留工作表、类型和公式；不把解析视图冒充完整文档。"""
import hashlib

import pytest
from openpyxl import Workbook

from src.agentic_runtime.candidate_qa import inspect_candidates
from src.agentic_runtime.candidate_verifier import CandidateVerifier, _candidate_preview
from src.agentic_runtime.models import PiRuntimeRequest, SemanticDecision, SourceInput


def workbook_file(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "本期"
    sheet.append(["编号", "金额", "启用"])
    sheet.append(["001", 12.5, True])
    sheet.append(["002", None, False])
    sheet = workbook.create_sheet("历史")
    sheet.append(["编号", "金额", "启用"])
    sheet.append(["003", "=SUM(1,2)", False])
    workbook.save(path)
    workbook.close()


def test_xlsx_candidate_keeps_sheet_type_and_formula_identity(tmp_path):
    workbook_file(tmp_path / "result.xlsx")
    preview = _candidate_preview(inspect_candidates(tmp_path, ("xlsx",))[0])
    assert 'WORKBOOK_SHEETS=["本期", "历史"]' in preview
    assert 'SHEET="本期"' in preview and 'SHEET="历史"' in preview
    assert '["001",12.5,true]' in preview and '["002",null,false]' in preview
    assert '{"formula":"=SUM(1,2)"}' in preview
    assert "NONEMPTY_DATA_ROWS=2" in preview
    assert "CELL_SCAN_COMPLETE=true" in preview
    assert "PARTIAL_TABLE_VIEW" in preview and "公式未计算" in preview


@pytest.mark.asyncio
@pytest.mark.parametrize("source_changed", [False, True])
async def test_semantic_check_receives_frozen_bounded_table_cells(tmp_path, source_changed):
    source = tmp_path / "source.xlsx"
    workbook_file(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    request = PiRuntimeRequest(user_id="owner", task_id="task", revision=1,
        objective_text="只输出本期的两行，保留编号字符串、金额数值与空值；不使用历史工作表。",
        model="fixture", base_url="http://127.0.0.1:1/v1", api_key="synthetic",
        requested_output_formats=("csv",), sources=(SourceInput(upload_id="source", original_name=source.name,
            host_path=source, sha256=digest),))
    if source_changed:
        source.write_bytes(source.read_bytes() + b"changed")
    output = tmp_path / "output"
    output.mkdir()
    (output / "result.csv").write_text("编号,金额,启用\n001,12.5,true\n002,,false\n", encoding="utf-8")
    captured = {}

    class Judge:
        async def judge(self, **kwargs):
            captured.update(kwargs)
            return SemanticDecision(passed=True, contains_unrequested_content=False, reason="合成判定")

    await CandidateVerifier(semantic_judge=Judge())._verify_semantics(request=request,
        candidates=inspect_candidates(output, ("csv",)), checks=[], grounded_evidence=(), grounded_quotes=())
    context = "\n".join(captured["evidence"])
    assert ("PARTIAL_TABLE_SOURCE=" in context) is (not source_changed)
    assert "FULL_FROZEN_SOURCE=" not in context
    if not source_changed:
        assert 'SHEET="本期"' in context and 'SHEET="历史"' in context
        assert '"001",12.5,true' in context


def test_large_xlsx_preview_is_bounded_and_discloses_omission(tmp_path):
    workbook = Workbook()
    workbook.active.append(["值"])
    for _ in range(600):
        workbook.active.append(["x" * 100])
    workbook.save(tmp_path / "large.xlsx")
    workbook.close()
    preview = _candidate_preview(inspect_candidates(tmp_path, ("xlsx",))[0])
    assert "TRUNCATED" in preview
    assert len(preview) < 21000
    assert "未读取部分" in preview


def test_sparse_wide_sheet_cannot_claim_complete_scan(tmp_path):
    workbook = Workbook()
    workbook.active.append(["值"])
    workbook.active.append([1])
    workbook.active.cell(3, 102, 99)
    workbook.save(tmp_path / "wide.xlsx")
    workbook.close()
    preview = _candidate_preview(inspect_candidates(tmp_path, ("xlsx",))[0])
    assert "TRUNCATED" in preview
    assert "CELL_SCAN_COMPLETE=true" not in preview


@pytest.mark.parametrize("formula_type", ["array", "data_table"])
def test_special_formula_is_literal_or_explicitly_unread(tmp_path, formula_type):
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
    workbook = Workbook()
    workbook.active.append(["值"])
    workbook.active["A2"] = (ArrayFormula(ref="A2", text="=SUM(1,2)") if formula_type == "array"
                              else DataTableFormula(ref="A2", r1="A1"))
    workbook.save(tmp_path / "formula.xlsx")
    workbook.close()
    preview = _candidate_preview(inspect_candidates(tmp_path, ("xlsx",))[0])
    assert "object at" not in preview
    if formula_type == "array":
        assert '{"formula":"=SUM(1,2)"}' in preview
    else:
        assert '"formula_unread":true' in preview and "CELL_SCAN_COMPLETE=false" in preview


@pytest.mark.parametrize("expected", [True, False])
def test_inconclusive_content_review_never_counts_as_success(expected):
    from scripts.evaluate_candidate_content import matches_expected
    assert not matches_expected("inconclusive", expected)
    assert not matches_expected("inflight_unknown", expected)
    assert matches_expected("passed" if expected else "failed", expected)


def test_content_summary_preserves_unknown_and_unrun_cases():
    from scripts.evaluate_candidate_content import summary
    assert summary([1, 2, 3], [
        {"status": "failed", "expected_passed": False},
        {"status": "inflight_unknown", "expected_passed": False},
    ]) == {"planned": 3, "completed": 2, "matched": 1, "unknown": 1, "not_run": 1}
