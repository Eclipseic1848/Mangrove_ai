"""只读工作簿窗口；与任务样本解析分开，保留物理单元格坐标。"""
from datetime import date, datetime, time
from pathlib import Path
import re
import threading
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

_SLOTS = threading.BoundedSemaphore(2)


class WorkbookPreviewError(ValueError):
    """允许向用户展示的预览限制，不包含解析器或宿主路径。"""


def _text(value, number_format=""):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S" if isinstance(value, datetime) and value.time() != time() else "%Y-%m-%d")
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        pattern = number_format.split(";")[0]
        match = re.search(r"\.([0#]+)", pattern)
        decimals = min(len(match.group(1)), 12) if match else 0
        if "%" in pattern:
            return f"{value * 100:.{decimals}f}%"
        if re.fullmatch(r"0+", pattern) and isinstance(value, int):
            return str(value).zfill(len(pattern))
        if match or "#,##" in pattern:
            return format(value, f",.{decimals}f" if "#,##" in pattern else f".{decimals}f")
    return str(value)


def _color(color):
    if color is not None and color.type == "rgb" and isinstance(color.rgb, str):
        return "#" + color.rgb[-6:]
    return None


def workbook_preview(path: Path, extension: str, sheet: int, row: int, column: int):
    if not _SLOTS.acquire(blocking=False):
        raise WorkbookPreviewError("预览服务忙，请稍后重试")
    try:
        return _read_workbook(path, extension, sheet, row, column)
    finally:
        _SLOTS.release()


def _read_workbook(path: Path, extension: str, sheet: int, row: int, column: int):
    if path.stat().st_size > 20 * 1024 * 1024:
        raise WorkbookPreviewError("工作簿超过 20 MB 在线预览上限，请下载原件查看")
    if extension == ".xls":
        return _xls_preview(path, sheet, row, column)
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(x.file_size for x in entries) > 32 * 1024 * 1024:
            raise WorkbookPreviewError("工作簿解压后超过安全预览上限，请下载原件查看")
        allocated = 0
        for entry in entries:
            if not entry.filename.startswith("xl/worksheets/") or not entry.filename.endswith(".xml"):
                continue
            raw = archive.read(entry)
            if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                raise WorkbookPreviewError("工作簿包含不允许的 XML 定义")
            for element in ElementTree.fromstring(raw).iter():
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "c":
                    allocated += 1
                elif tag == "mergeCell":
                    left, top, right, bottom = range_boundaries(element.attrib["ref"])
                    allocated += (right-left+1) * (bottom-top+1)
                    if allocated > 100000:
                        raise WorkbookPreviewError("合并范围超过十万个单元格预览上限，请下载原件查看")
                if allocated > 100000:
                    raise WorkbookPreviewError("工作簿超过十万个单元格预览上限，请下载原件查看")
    # 不更新外部链接，不运行公式；只展示文件保存的公式结果。
    with path.open("rb") as stream:
        book = load_workbook(stream, data_only=True, keep_links=False)
    with path.open("rb") as stream:
        formulas = load_workbook(stream, data_only=False, keep_links=False)
    try:
        if not 0 <= sheet < len(book.worksheets):
            raise WorkbookPreviewError("工作表不存在")
        ws = book.worksheets[sheet]
        end_row, end_column = min(row + 100, ws.max_row), min(column + 30, ws.max_column)
        cells = []
        def cell_text(cell):
            formula = formulas.worksheets[sheet].cell(cell.row, cell.column)
            return f"{formula.value}（结果未缓存）" if cell.value is None and formula.data_type == "f" else _text(cell.value, cell.number_format)
        for r in range(row + 1, end_row + 1):
            line = []
            for c in range(column + 1, end_column + 1):
                cell = ws.cell(r, c)
                style = {"fontWeight": "bold" if cell.font.bold else "normal"}
                if cell.font.italic:
                    style["fontStyle"] = "italic"
                if cell.alignment.horizontal in {"left", "center", "right"}:
                    style["textAlign"] = cell.alignment.horizontal
                if _color(cell.font.color):
                    style["color"] = _color(cell.font.color)
                if cell.fill.patternType == "solid" and _color(cell.fill.fgColor):
                    style["backgroundColor"] = _color(cell.fill.fgColor)
                line.append({"text": cell_text(cell), "style": style})
            cells.append(line)
        merges = []
        for merged in ws.merged_cells.ranges:
            if merged.min_row <= end_row and merged.max_row > row and merged.min_col <= end_column and merged.max_col > column:
                # 跨窗口合并格携带锚点正文，避免进入下一窗口后看见空白。
                merges.append({"row": merged.min_row - 1, "column": merged.min_col - 1,
                               "rows": merged.max_row - merged.min_row + 1, "columns": merged.max_col - merged.min_col + 1})
                top, left = max(merged.min_row - 1, row), max(merged.min_col - 1, column)
                if top < end_row and left < end_column:
                    anchor = ws.cell(merged.min_row, merged.min_col)
                    cells[top - row][left - column]["text"] = cell_text(anchor)
        return {"sheets": book.sheetnames, "sheet": sheet, "row": row, "column": column,
                "total_rows": ws.max_row, "total_columns": ws.max_column, "cells": cells, "merges": merges,
                "widths": [max(60, min(600, (ws.column_dimensions[get_column_letter(c)].width or 13) * 7 + 8)) for c in range(column + 1, end_column + 1)],
                "heights": [max(28, min(400, (ws.row_dimensions[r].height or 21) * 4 / 3)) for r in range(row + 1, end_row + 1)],
                "note": "按窗口浏览原工作表；公式显示已保存结果，复杂数字格式、图表和条件格式可能与 Excel 不同。"}
    finally:
        book.close()
        formulas.close()


def _xls_preview(path, sheet, row, column):
    import xlrd
    book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    try:
        if not 0 <= sheet < book.nsheets:
            raise WorkbookPreviewError("工作表不存在")
        ws = book.sheet_by_index(sheet)
        if ws.nrows * ws.ncols > 100000:
            raise WorkbookPreviewError("旧版工作簿超过十万个单元格预览上限")
        cells = []
        for r in range(row, min(row + 100, ws.nrows)):
            line = []
            for c in range(column, min(column + 30, ws.ncols)):
                cell = ws.cell(r, c)
                value = xlrd.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                xf = book.xf_list[cell.xf_index]
                fmt = book.format_map.get(xf.format_key)
                line.append({"text": _text(value, fmt.format_str if fmt else ""),
                             "style": {"fontWeight": "bold" if book.font_list[xf.font_index].bold else "normal"}})
            cells.append(line)
        for a, b, c, d in ws.merged_cells:
            top, left = max(a, row), max(c, column)
            if top < min(b, row+len(cells)) and left < min(d, column+(len(cells[0]) if cells else 0)):
                cells[top-row][left-column]["text"] = _text(ws.cell_value(a, c))
        return {"sheets": book.sheet_names(), "sheet": sheet, "row": row, "column": column,
                "total_rows": ws.nrows, "total_columns": ws.ncols, "cells": cells,
                "merges": [{"row": a, "column": c, "rows": b-a, "columns": d-c} for a,b,c,d in ws.merged_cells],
                "widths": [120] * min(30, max(0, ws.ncols-column)), "heights": [28] * len(cells),
                "note": "旧版 XLS 只读预览；复杂样式请下载原件核对。"}
    finally:
        book.release_resources()
