# -*- coding: utf-8 -*-
"""Pi 候选文件的确定性完整性检查。"""
from __future__ import annotations

import csv
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path

from .models import CandidateArtifact


_EXTENSION_FORMAT = {".md": "markdown"}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _reopen(path: Path, fmt: str) -> tuple[str, ...]:
    checks = ["non_empty"]
    if fmt == "json":
        json.loads(path.read_text(encoding="utf-8"))
    elif fmt == "jsonl":
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                json.loads(line)
    elif fmt == "csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        if not rows or not any(cell.strip() for cell in rows[0]):
            raise ValueError("CSV 缺少表头")
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if not workbook.worksheets:
                raise ValueError("XLSX 没有工作表")
        finally:
            workbook.close()
    elif fmt == "parquet":
        import pyarrow.parquet as pq

        pq.read_metadata(path)
    elif fmt == "docx":
        from docx import Document

        Document(path)
    elif fmt == "pdf":
        from pypdf import PdfReader

        if not PdfReader(path).pages:
            raise ValueError("PDF 没有页面")
    elif fmt == "pptx":
        from pptx import Presentation

        if not Presentation(path).slides:
            raise ValueError("PPTX 没有页面")
    elif fmt == "html":
        parser = HTMLParser()
        parser.feed(path.read_text(encoding="utf-8"))
        parser.close()
    elif fmt in {"txt", "markdown"}:
        if not path.read_text(encoding="utf-8-sig").strip():
            raise ValueError(f"{fmt.upper()} 输出为空")
    else:
        raise ValueError(f"不支持检查候选格式：{fmt}")
    return (*checks, "reopened")


def inspect_candidates(
    output_dir: Path,
    requested_formats: tuple[str, ...],
) -> tuple[CandidateArtifact, ...]:
    """只登记请求格式且能重新打开的普通文件。

    这里故意拒绝符号链接和目录逃逸：即使容器内 Agent 能自由执行命令，也不能借候选
    下载接口把任务目录之外的宿主文件带出。
    """

    root = output_dir.resolve()
    normalized = {
        "markdown" if item == "md" else item
        for item in requested_formats
    }
    candidates: list[CandidateArtifact] = []
    for path in sorted(output_dir.iterdir()):
        if (
            not path.is_file()
            or path.is_symlink()
            or path.name == "candidate-manifest.json"
        ):
            continue
        resolved = path.resolve()
        if root not in resolved.parents:
            raise ValueError("候选文件越过任务输出目录")
        fmt = _EXTENSION_FORMAT.get(
            path.suffix.lower(),
            path.suffix.lower().lstrip("."),
        )
        if fmt not in normalized:
            continue
        if path.stat().st_size <= 0:
            raise ValueError(f"候选文件为空：{path.name}")
        checks = _reopen(path, fmt)
        digest = _sha256(path)
        candidates.append(
            CandidateArtifact(
                artifact_id=f"candidate_{digest[:16]}",
                filename=path.name,
                format=fmt,
                host_path=resolved,
                sha256=digest,
                size_bytes=path.stat().st_size,
                openable=True,
                qa_checks=checks,
            )
        )
    if not candidates:
        requested = "、".join(sorted(normalized))
        raise ValueError(f"Pi 未生成可重新打开的请求格式文件：{requested}")
    return tuple(candidates)
