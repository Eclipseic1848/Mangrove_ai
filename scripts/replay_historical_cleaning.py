"""隔离重放历史清洗与文档交付节点；历史产物仅作回归基线，不证明业务正确。"""
from __future__ import annotations

import argparse
import asyncio
from collections import Counter
from contextlib import closing
import hashlib
from io import BytesIO
import json
from pathlib import Path
import socket
import sqlite3
import tempfile
import time
from unittest.mock import patch

from scripts.audit_historical_task_contracts import write_report
from src.data_prep.artifact_store import ArtifactStore
from src.data_prep.models import DataPrepTaskSpec


def _digest(data):
    return hashlib.sha256(data).hexdigest()


def _rows(data):
    return [json.loads(line) for line in data.decode("utf-8").splitlines() if line.strip()]


def _xlsx_cells(data):
    from openpyxl import load_workbook

    # 比较表名、顺序、值与类型；保留公式文本，不依赖可能过期的公式缓存或ZIP时间戳。
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        return [(sheet.title, tuple(tuple((cell.data_type, cell.value) for cell in row)
                                   for row in sheet.iter_rows())) for sheet in workbook]
    finally:
        workbook.close()


def _replay_document(task_id, payload, manifest, read_local, entry):
    from src.data_prep import document_models as models
    from src.data_prep.models import RawArtifact
    from src.services.document_delivery import write_document_delivery

    if manifest.get("task_id") != task_id:
        raise ValueError("Manifest不属于本任务")
    indexed = {item["kind"]: item for item in manifest["artifacts"] if item["kind"] != "raw"}
    required = ("extraction_spec", "extracted_fields_json", "extracted_records", "extracted_tables",
                "extracted_documents", "review_tasks", "quality", "evidence", "lineage", "schema")
    missing = [kind for kind in required if kind not in indexed]
    if missing:
        # 一次列全清单中缺少的快照；不能把旧版本未登记的输入直接当成空值。
        entry.update(missing_snapshot_kind=missing[0], missing_snapshot_kinds=missing)
        raise ValueError("缺少抽取快照，不能假定为空")
    observed = {}

    def read_item(item):
        data = read_local(item["path"])
        digest = _digest(data)
        if digest != item["sha256"]:
            raise ValueError("历史制品哈希不匹配")
        observed[item["path"]] = digest
        return data

    def values(kind, model):
        item = indexed[kind]
        data = read_item(item)
        rows = _rows(data) if item["path"].endswith(".jsonl") else json.loads(data)
        return [model.model_validate(row) for row in rows]

    spec = models.ExtractionSpec.model_validate_json(read_item(indexed["extraction_spec"]))
    arguments = {"fields": values("extracted_fields_json", models.ExtractedField),
                 "records": values("extracted_records", models.ExtractedRecord),
                 "tables": values("extracted_tables", models.ExtractedTable),
                 "documents": values("extracted_documents", models.ExtractedDocument),
                 "review_tasks": values("review_tasks", models.ReviewTask)}
    arguments["parse_rejects"] = json.loads(read_item(indexed["rejects"])) if "rejects" in indexed else []
    if "extraction_coverage" in indexed:
        arguments["coverage"] = json.loads(read_item(indexed["extraction_coverage"]))
    raw = [RawArtifact.model_validate(item) for item in payload["raw_artifacts"]]
    if any(item.task_id != task_id for item in raw):
        raise ValueError("来源不属于历史任务")
    expected_quality = json.loads(read_item(indexed["quality"]))["overall"]
    # 对比权威结构化输出、证据和血缘；不把重新序列化当作模型或OCR重放。
    comparison = [item for item in manifest["outputs"] if item["format"] in {"json", "jsonl", "xlsx"}]
    comparison += [indexed[kind] for kind in ("evidence", "lineage", "schema")]
    baselines = {item["path"]: read_item(item) for item in comparison}
    with tempfile.TemporaryDirectory(prefix="mangrove-history-document-") as temporary:
        target = ArtifactStore(temporary)
        with patch.object(socket.socket, "connect", side_effect=PermissionError("离线重放禁止网络")), patch.object(socket.socket, "connect_ex", side_effect=PermissionError("离线重放禁止网络")):
            result = write_document_delivery(target, task_id, spec=spec, raw_artifacts=raw, **arguments)
        differences = []
        for name, data in baselines.items():
            parser = _xlsx_cells if name.endswith(".xlsx") else _rows if name.endswith(".jsonl") else json.loads
            generated = target.resolve_path(name)
            if not generated.is_file() or parser(generated.read_bytes()) != parser(data):
                differences.append(Path(name).name)
        current_manifest = json.loads(target.resolve_path(result.manifest_path).read_bytes())
        if [(item["format"], item["records"]) for item in current_manifest["outputs"]] != [(item["format"], item["records"]) for item in manifest["outputs"]]:
            differences.append("output_contract")
        if result.quality.overall.value != expected_quality:
            differences.append("quality_overall")
    entry.update(stage="extraction_snapshot_to_delivery_and_quality", status="failed" if differences else "passed",
                 reason="historical_document_results_differ" if differences else "matches_historical_document_results",
                 differences=differences, compared_artifacts=len(comparison), input_sha256=observed,
                 source_unchanged=all(_digest(read_local(name)) == digest for name, digest in observed.items()))
    if not entry["source_unchanged"]:
        entry.update(status="blocked", reason="historical_source_changed_during_replay")


async def replay(database: Path, artifacts: Path) -> dict:
    from src.data_prep import graph

    source = ArtifactStore(str(artifacts))
    with closing(sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)) as connection:
        tasks = connection.execute("SELECT task_id,spec_json,manifest_path FROM data_prep_tasks ORDER BY task_id").fetchall()
    entries = []
    for task_id, raw_spec, manifest_path in tasks:
        started = time.perf_counter()
        entry = {"task_id": task_id, "status": "blocked", "business_correctness": "unverified"}
        try:
            if Path(task_id).name != task_id or task_id in {".", ".."}:
                raise ValueError("任务标识不能作为目录路径")
            payload = json.loads(raw_spec)
            entry["spec_sha256"] = _digest(raw_spec.encode("utf-8"))
            task_root = source.resolve_path(task_id)

            def read_local(relative):
                path = source.resolve_path(relative)
                if not path.is_relative_to(task_root):
                    raise ValueError("制品不属于本任务")
                return path.read_bytes()

            if not manifest_path:
                entry["reason"] = "missing_manifest"
                continue
            manifest_bytes = read_local(manifest_path)
            manifest = json.loads(manifest_bytes)
            if payload.get("task_type") == "document_extraction":
                _replay_document(task_id, payload, manifest, read_local, entry)
                if read_local(manifest_path) != manifest_bytes:
                    entry.update(status="blocked", reason="historical_source_changed_during_replay")
                continue
            if payload.get("task_type"):
                entry["reason"] = "unsupported_historical_task_type"
                continue
            spec = DataPrepTaskSpec.model_validate(payload)
            expected = next((item for item in manifest.get("outputs", []) if item["format"] == "jsonl"), None)
            if expected is None:
                entry["reason"] = "missing_historical_jsonl_baseline"
                continue
            golden = read_local(expected["path"])
            if _digest(golden) != expected["sha256"]:
                entry["reason"] = "historical_output_hash_mismatch"
                continue
            expected_rows = _rows(golden)
            batches = sorted((task_root / "parsed").glob("part-*.jsonl"))
            if not batches:
                entry["reason"] = "missing_parsed_snapshot"
                continue
            entry["manifest_sha256"] = _digest(manifest_bytes)
            entry["baseline_sha256"] = _digest(golden)
            entry["parsed_sha256"] = []
            with tempfile.TemporaryDirectory(prefix="mangrove-history-clean-") as temporary:
                target = ArtifactStore(temporary)
                refs = []
                for number, path in enumerate(batches):
                    data = read_local(str(path))
                    entry["parsed_sha256"].append(_digest(data))
                    refs.append(target.append_jsonl_batch(task_id, "parsed", _rows(data), number))
                # 只替换制品根目录，执行真实节点；禁止采集、模型调用及任何网络连接。
                with patch.object(graph, "ArtifactStore", lambda: target), patch.object(socket.socket, "connect", side_effect=PermissionError("离线重放禁止网络")), patch.object(socket.socket, "connect_ex", side_effect=PermissionError("离线重放禁止网络")):
                    result = await graph.clean_node({"task_id": task_id, "spec": spec, "parsed_batches": refs})
                actual = [row for ref in result["clean_batches"] for row in _rows(target.resolve_path(ref.path).read_bytes())]
                entry.update(status="passed" if actual == expected_rows else "failed",
                             reason="matches_historical_rows" if actual == expected_rows else "historical_rows_differ",
                             actual_records=len(actual), expected_records=len(expected_rows))
            entry["source_unchanged"] = (
                _digest(read_local(manifest_path)) == entry["manifest_sha256"]
                and _digest(read_local(expected["path"])) == entry["baseline_sha256"]
                and [_digest(read_local(str(path))) for path in batches] == entry["parsed_sha256"]
            )
            if not entry["source_unchanged"]:
                entry.update(status="blocked", reason="historical_source_changed_during_replay")
        except Exception as error:
            # 报告不包含异常正文，避免历史业务内容或宿主路径进入可共享摘要。
            entry.update(status="blocked", reason="replay_error", error_type=type(error).__name__)
        finally:
            entry["duration_ms"] = round((time.perf_counter() - started) * 1000, 3)
            entries.append(entry)
    return {"stage": "historical_snapshot_node_replay", "entries": entries,
            "counts": dict(Counter(item["status"] for item in entries)), "business_success_rate": None,
            "limitations": ["不重放采集、解析、OCR、模型或正式发布", "文档仅重放抽取快照到序列化/证据/血缘/质量；XLSX比较表名、顺序、单元格值和类型，不检查样式或计算公式", "历史产物不是独立人工正确性标注", "解析快照哈希为本次读取记录，不是历史冻结证明"]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", type=Path, required=True)
    parser.add_argument("--artifacts", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.resolve().is_relative_to(args.artifacts.resolve()):
        raise ValueError("评测报告不能覆盖历史制品目录")
    if args.output.exists():
        raise FileExistsError("重放报告必须使用新文件名")
    report = asyncio.run(replay(args.database, args.artifacts))
    write_report(args.database, args.output, report, overwrite=False)
    print(json.dumps(report["counts"]))
    return int(any(item["status"] != "passed" for item in report["entries"]))


if __name__ == "__main__":
    raise SystemExit(main())
