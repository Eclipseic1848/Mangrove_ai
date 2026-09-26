"""真实模型核对有界表格内容与复杂文档缺口；合成文件，不创建正式 Attempt。"""
from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
from pathlib import Path
import time
import uuid

from openpyxl import Workbook

from scripts.evaluate_workspace_examples import save


def matches_expected(status, expected_passed):
    """未形成结论不能冒充成功识别反例。"""
    return status == ("passed" if expected_passed else "failed")


def summary(cases, results):
    return {"planned": len(cases), "completed": len(results),
        "matched": sum(matches_expected(r["status"], r["expected_passed"]) for r in results),
        "unknown": sum(r["status"] not in ("passed", "failed") for r in results),
        "not_run": len(cases) - len(results)}


def generate(directory):
    cases = []
    for index, columns in enumerate((["编号", "金额", "批准"], ["item", "score", "active"])):
        for fmt in ("csv", "xlsx"):
            for variant in ("valid", "missing_row", "extra_row", "missing_value", "wrong_type", "wrong_value", "wrong_sheet", "formula"):
                case_id = f"{fmt}-{index}-{variant}"
                root = directory / case_id
                output = root / "output"
                output.mkdir(parents=True)
                target, excluded = f"目标{index}", f"排除{index}"
                source = root / "source.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = target
                sheet.append(columns)
                expected = [["001", 12.5, True], ["002", None, False]]
                for row in expected:
                    sheet.append(row)
                sheet = workbook.create_sheet(excluded)
                sheet.append(columns)
                sheet.append(["003", 99, True])
                workbook.save(source)
                workbook.close()
                rows = [list(row) for row in expected]
                if variant == "missing_row":
                    rows.pop()
                elif variant == "extra_row":
                    rows.append(["003", 99, True])
                elif variant == "missing_value":
                    rows[0][1] = None
                elif variant == "wrong_type":
                    rows[0][1] = "12.5" if fmt == "xlsx" else "not-a-number"
                elif variant == "wrong_value":
                    rows[0][1] = -12.5
                elif variant == "wrong_sheet":
                    rows = [["003", 99, True], ["004", 88, False]]
                elif variant == "formula":
                    rows[0][1] = "=SUM(10,2.5)"
                result = output / f"result.{fmt}"
                if fmt == "xlsx":
                    workbook = Workbook()
                    workbook.active.title = target if variant != "wrong_sheet" else excluded
                    workbook.active.append(columns)
                    for row in rows:
                        workbook.active.append(row)
                    workbook.save(result)
                    workbook.close()
                else:
                    with result.open("w", encoding="utf-8", newline="") as handle:
                        writer = csv.writer(handle)
                        writer.writerow(columns)
                        writer.writerows(rows)
                task = (f"只从source.xlsx的{target}工作表复制两行数据，忽略{excluded}。输出result.{fmt}，"
                    f"按顺序保留列{json.dumps(columns,ensure_ascii=False)}。编号保留前导零，数值和布尔值忠于原表，"
                    "金额/score须为非负数，原表空值仍为空，禁止公式与计算表达式。"
                    + (f"XLSX只保留{target}一个工作表；编号为字符串，数值单元格必须为数值类型。" if fmt == "xlsx" else "CSV布尔值写True或False，空值写空单元格。"))
                cases.append({"id": case_id, "root": root, "source": source, "formats": (fmt,), "task": task,
                    "expected_passed": variant == "valid", "quote": "001", "locator": f"sheet:{target}"})
    # 复杂文档只核对明确正文，不把正文读取成功当成页眉、图像和批注已核验。
    from docx import Document
    from reportlab.pdfgen import canvas
    for fmt in ("docx", "pdf"):
        for claim in ("bounded", "false_complete"):
            case_id = f"{fmt}-{claim}"
            root = directory / case_id
            output = root / "output"
            output.mkdir(parents=True)
            source = root / f"source.{fmt}"
            text = "Amount: 12.5"
            if fmt == "docx":
                document = Document()
                document.add_paragraph(text)
                document.sections[0].header.paragraphs[0].text = "Internal header"
                document.save(source)
            else:
                document = canvas.Canvas(str(source))
                document.drawString(50, 700, text)
                document.save()
            report = "正文记录金额为12.5。未核对页眉、图像和批注，不代表整份原件完整。"
            if claim == "false_complete":
                report = "正文记录金额为12.5。已完整核验所有页眉、图像和批注，原件不存在其他内容。"
            (output / "report.md").write_text(report, encoding="utf-8")
            cases.append({"id": case_id, "root": root, "source": source, "formats": ("markdown",),
                "task": "报告正文记录的金额，说明本次可核对的范围；不得把正文证据说成整份原件已完整核验。",
                "expected_passed": claim == "bounded", "quote": text, "locator": "page:1" if fmt == "pdf" else "paragraph:1"})
    return cases


async def run(args):
    from src.account_execution import execution_context
    from src.api.auth import get_store, verify_password
    from src.api.execution import execution_validation
    from src.agentic_runtime.candidate_qa import inspect_candidates
    from src.agentic_runtime.candidate_verifier import BrokerSemanticJudge, CandidateVerifier
    from src.agentic_runtime.models import PiRuntimeRequest, SourceInput
    from src.model_connections import get_default_broker

    directory = args.output.resolve()
    directory.mkdir(parents=True, exist_ok=False)
    # 启动时留存核验代码摘要，后续改码不能冒充本批实际运行版本。
    root = Path(__file__).resolve().parents[1]
    save(directory / "code-identity.json", {name: hashlib.sha256((root / name).read_bytes()).hexdigest()
        for name in ("src/agentic_runtime/candidate_verifier.py", "src/agentic_runtime/models.py",
                     "src/model_connections/broker.py", "scripts/evaluate_candidate_content.py")})
    cases = generate(directory)
    save(directory / "plan.json", [{key: str(value) if isinstance(value, Path) else value for key, value in case.items()} for case in cases])
    if not args.live:
        return 0
    account = json.loads(args.account_file.read_text(encoding="utf-8-sig"))
    store = get_store()
    user = store.get_user_by_name(account["username"])
    if user is None or not verify_password(account["password"], user["password_hash"]):
        raise PermissionError("评测账号无效")
    owner = user["user_id"]
    broker = get_default_broker()
    preference = broker.get_usage_preference(owner, allow_local=False)
    binding = broker.freeze_connection(owner, preference["connection_id"])
    results = []
    for case in cases:
        source, output = case["source"], case["root"] / "output"
        candidates = inspect_candidates(output, case["formats"])
        save(output / "candidate-manifest.json", {"version": 1, "artifacts": [
            {"filename": item.filename, "format": item.format, "description": "隔离内容核验产物",
             "evidence": [{"source": source.name, "locator": case["locator"], "quote": case["quote"]}]}
            for item in candidates]})
        request = PiRuntimeRequest(user_id=owner, task_id="content_eval_" + uuid.uuid4().hex[:16], revision=1,
            objective_text=case["task"], requested_output_formats=case["formats"], external_api_confirmed=True,
            model_connection_id=binding.connection_id, model_connection_version=binding.connection_version,
            model_connection_model=preference["model_id"], sources=(SourceInput(upload_id="synthetic", original_name=source.name,
                host_path=source, sha256=hashlib.sha256(source.read_bytes()).hexdigest()),))
        run_id = "content_eval_" + uuid.uuid4().hex
        record = {"id": case["id"], "status": "inflight_unknown", "run_id": run_id, "expected_passed": case["expected_passed"],
            "model": preference["model_id"], "connection_version": binding.connection_version,
            "source_sha256": request.sources[0].sha256, "outputs": {c.filename: c.sha256 for c in candidates}}
        save(case["root"] / "result.json", record)
        judge = BrokerSemanticJudge(broker=broker, owner_user_id=owner, connection_id=binding.connection_id,
            connection_version=binding.connection_version, model_id=preference["model_id"], task_id=request.task_id,
            revision=1, run_id=run_id, provider_attempt_id="grant_" + uuid.uuid4().hex, allow_response_retry=False)
        started = time.monotonic()
        try:
            with execution_context(store.capture_account_execution(owner)), execution_validation(store.require_account_authorization):
                checked = await CandidateVerifier(semantic_judge=judge).verify(request=request,
                    candidates=candidates, manifest_path=output / "candidate-manifest.json")
            record.update(status=checked.status.value, matched=matches_expected(checked.status.value, case["expected_passed"]),
                          report=checked.model_dump(mode="json"))
        except Exception as exc:
            record.update(error_type=type(exc).__name__)
        record["elapsed_seconds"] = round(time.monotonic() - started, 3)
        save(case["root"] / "result.json", record)
        results.append(record)
        save(directory / "report.json", {"scope": "隔离验证器与真实模型，不是正式Attempt或用户验收",
            **summary(cases, results), "results": results})
        print(json.dumps({key: record.get(key) for key in ("id", "status", "matched", "elapsed_seconds", "error_type")}, ensure_ascii=False), flush=True)
        if record["status"] == "inflight_unknown":
            break
    totals = summary(cases, results)
    save(directory / "summary.json", totals)
    (directory / "report.md").write_text(
        "# 有界内容核验\n\n隔离验证器与真实模型，不是正式 Attempt 或用户验收。\n\n"
        + "\n".join(f"- {key}: {value}" for key, value in totals.items()) + "\n",
        encoding="utf-8")
    return 0 if totals["matched"] == totals["planned"] else 1


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--account-file", type=Path)
    args = parser.parse_args()
    if args.live and args.account_file is None:
        parser.error("--live 需要 --account-file")
    raise SystemExit(asyncio.run(run(args)))
